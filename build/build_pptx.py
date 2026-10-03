# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from data import *

H = lambda s: RGBColor.from_string(s.lstrip("#"))
NAVY_C, NAVY2_C, ACC_C = H("14264B"), H("22386B"), H("F5A800")
ACC_L, BLUE_L, GREY_L, INK_C, MUTED_C = H("FFF1CC"), H("E8EEF9"), H("EDEFF4"), H("1A2238"), H("5B6680")
WHITE = H("FFFFFF")
ANS = {"A": (H("D62839"), WHITE), "B": (H("1F6FD1"), WHITE), "C": (H("F7C600"), NAVY_C), "D": (H("2E9E4F"), WHITE)}
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
IMG = os.path.join(os.path.dirname(__file__), "img")


def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb = line; r.line.width = Pt(1)
    r.shadow.inherit = False
    return r


def text(s, x, y, w, h, t, size=20, bold=False, color=INK_C, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None):
    tb = shape or s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05)
    paras = t if isinstance(t, list) else [t]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = ptxt if isinstance(ptxt, list) else [(ptxt, bold, color)]
        for rt, rb, rc in runs:
            r = p.add_run(); r.text = rt; r.font.size = Pt(size); r.font.bold = rb; r.font.color.rgb = rc; r.font.name = FONT
        if i: p.space_before = Pt(size * 0.5)
    return tb


def slide(title=None, kicker=None, notes=None, dark=False):
    s = prs.slides.add_slide(BLANK)
    if dark:
        rect(s, 0, 0, 13.333, 7.5, NAVY_C); rect(s, 0, 0, 0.35, 7.5, ACC_C)
    else:
        rect(s, 0, 0, 0.35, 7.5, ACC_C)
        if kicker: text(s, 0.7, 0.3, 11, 0.35, kicker.upper(), 13, True, H("B07A00"))
        if title: text(s, 0.7, 0.6, 12, 0.9, title, 32, True, NAVY_C)
    if notes: s.notes_slide.notes_text_frame.text = notes
    return s


def table(s, data, x, y, w, colw, size=16, rowh=0.5, bold_first=True, head=True):
    rows, cols = len(data), len(data[0])
    gt = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(rowh * rows))
    t = gt.table
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(data):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            c.margin_left = c.margin_right = Inches(0.1); c.margin_top = c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = val; r.font.name = FONT; r.font.size = Pt(size)
            c.text_frame.word_wrap = True
            c.fill.solid()
            if i == 0 and head:
                c.fill.fore_color.rgb = NAVY_C; r.font.color.rgb = WHITE; r.font.bold = True
            else:
                c.fill.fore_color.rgb = H("F6F8FC") if i % 2 == 0 else WHITE
                r.font.color.rgb = INK_C
                if j == 0 and bold_first: r.font.bold = True; r.font.color.rgb = NAVY_C
    return t


def card(s, x, y, w, h, head, body, hsize=20, bsize=18, fill=BLUE_L, bar=NAVY2_C, off=None):
    rect(s, x, y, w, h, fill); rect(s, x, y, 0.1, h, bar)
    text(s, x + 0.25, y + 0.08, w - 0.4, 0.5, head, hsize, True, NAVY_C)
    off = off or 0.12 + hsize * 0.026
    text(s, x + 0.25, y + off, w - 0.4, h - off - 0.05, body, bsize)


def question(station, label, q, answers, extra=None, notes=None):
    s = slide(notes=notes)
    rect(s, 0.35, 0, 12.983, 0.7, NAVY_C)
    text(s, 0.7, 0.1, 9, 0.5, f"{station} · {label}", 18, True, ACC_C, anchor=MSO_ANCHOR.MIDDLE)
    if extra: text(s, 9.2, 0.1, 3.8, 0.5, extra, 20, True, WHITE, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    size = 32 if len(q) < 110 else 26
    text(s, 0.7, 0.95, 12.1, 2.35, q, size, True, NAVY_C, anchor=MSO_ANCHOR.MIDDLE)
    asize = 24 if max(len(a) for a in answers) < 60 else 20
    for i, (L, a) in enumerate(zip("ABCD", answers)):
        cx = 0.7 + (i % 2) * 6.15; cy = 3.45 + (i // 2) * 1.95
        bg, fg = ANS[L]
        rect(s, cx, cy, 5.95, 1.8, bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
        c = rect(s, cx + 0.2, cy + 0.45, 0.9, 0.9, WHITE, shape=MSO_SHAPE.OVAL)
        text(s, 0, 0, 0, 0, L, 34, True, bg if L != "C" else NAVY_C, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, shape=c)
        text(s, cx + 1.2, cy + 0.1, 4.6, 1.6, a, asize, True, fg, anchor=MSO_ANCHOR.MIDDLE)
    return s


def intro(station, title, steps, notes=None):
    s = slide(title, station, notes)
    y = 1.9
    for i, st in enumerate(steps):
        rect(s, 0.8, y + 0.05, 0.6, 0.6, NAVY_C, shape=MSO_SHAPE.OVAL)
        text(s, 0, 0, 0, 0, str(i + 1), 22, True, ACC_C, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, shape=s.shapes[-1])
        text(s, 1.65, y - 0.05, 10.8, 1.0, st, 26, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.2
    return s


# =========================================================== 1 Titel
s = slide(dark=True, notes="„Guten Tag zusammen, wir sind Yasin und Mido. Wir stellen euch unser Projekt ‚Rechte und Pflichten aus dem Ausbildungsvertrag‘ vor: kurz, wie wir gearbeitet haben, dann unser Prüfungsvorbereitungsbooklet – und danach trainiert ihr selbst mit Aufgaben und Quizspielen.“")
text(s, 1.0, 1.2, 11.5, 0.5, "PRÜFUNGSVORBEREITUNG · WIRTSCHAFTS- UND SOZIALKUNDE", 16, True, ACC_C)
text(s, 1.0, 1.8, 11.5, 2.0, "Rechte und Pflichten aus dem Ausbildungsvertrag", 48, True, WHITE)
text(s, 1.0, 3.9, 11, 0.5, "Yasin & Mido", 24, False, H("DCE3F2"))
for i, t in enumerate(["Projekt", "Booklet", "Prüfungstraining"]):
    x = 1.0 + i * 3.9
    rect(s, x, 5.2, 3.6, 1.1, NAVY2_C); rect(s, x, 5.2, 3.6, 0.08, ACC_C)
    text(s, x, 5.35, 3.6, 0.9, [[(f"{i+1}  ", True, ACC_C), (t, True, WHITE)]], 24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# =========================================================== 2 Projekt in Kuerze
s = slide("Unser Projekt in Kürze", "Aufgabe, Problem, Ziel",
          "„Unsere Aufgabe: ein Prüfungsvorbereitungsbooklet zu einem Thema der schriftlichen IHK-Abschlussprüfung in Wirtschafts- und Sozialkunde. Wir haben den Ausbildungsvertrag gewählt, weil er Prüfungsstoff ist und sich mit echten Beispielen gut erklären lässt. Bisher gab es dazu kein kompaktes Lernmaterial mit Aufgaben und Lösungen. Unser Ziel: Am 25. November sollt ihr die Regeln zu Beginn und Dauer, Arbeitszeit, Vergütung, Urlaub, Probezeit und Kündigung sicher anwenden können – und das üben wir gleich.“")
items = [("Aufgabe", "Ein Prüfungsvorbereitungsbooklet zu einem Thema der schriftlichen IHK-Abschlussprüfung in Wirtschafts- und Sozialkunde."),
         ("Problem", "Zum Ausbildungsvertrag gab es bisher kein kompaktes Lernmaterial mit Aufgaben und Lösungen."),
         ("Ziel", "Regeln zu Beginn und Dauer, Arbeitszeit, Vergütung, Urlaub, Probezeit und Kündigung sicher anwenden können.")]
for i, (h, b) in enumerate(items):
    card(s, 0.7, 1.7 + i * 1.7, 8.6, 1.55, h, b, 20, 18)
rect(s, 9.7, 1.75, 3.1, 4.35, NAVY_C); rect(s, 9.7, 1.75, 3.1, 0.1, ACC_C)
text(s, 9.7, 2.3, 3.1, 0.5, "PRÜFUNG", 18, True, ACC_C, PP_ALIGN.CENTER)
text(s, 9.7, 2.9, 3.1, 1.3, "25.11.", 54, True, WHITE, PP_ALIGN.CENTER)
text(s, 9.7, 4.4, 3.1, 1.4, "schriftliche Abschlussprüfung", 18, False, H("DCE3F2"), PP_ALIGN.CENTER)

# =========================================================== 3a Zeitplan
s = slide("Vorgehen und Zeitplan", "Durchführung · Soll", "„Für die Umsetzung hatten wir die Woche bis zum 8. Oktober und sind in fünf Schritten vorgegangen: Recherche und Vertragssichtung, Gliederung, Texte mit Praxisbeispielen, Layout mit Korrekturlesen und zum Schluss die Präsentation. Texte und Praxisbeispiele haben wir parallel erarbeitet, um Zeit zu sparen. Die Aufgaben waren so verteilt: Recherche und Texte bei Yasin, Vertragssichtung, Praxisbeispiele, Grafiken und Korrekturlesen bei Mido, Gliederung, Layout und Präsentation gemeinsam.“")
table(s, [["Schritt", "Soll", "Wer", "Ist"],
          ["Recherche und Vertragssichtung", "01.–02.10.", "Yasin (Recherche), Mido (Vertrag)", ""],
          ["Gliederung", "02.–03.10.", "Yasin & Mido", ""],
          ["Texte, Praxisbeispiele und Grafiken", "03.–05.10.", "Yasin (Texte), Mido (Beispiele, Grafiken)", ""],
          ["Layout und Korrekturlesen", "05.–06.10.", "Yasin & Mido, Mido (Korrekturlesen)", ""],
          ["Präsentation vorbereiten und proben", "06.–07.10.", "Yasin & Mido", ""],
          ["Präsentation", "08.10.", "Yasin & Mido", ""]],
      0.7, 1.8, 12.0, [4.6, 1.7, 4.2, 1.5], 18, 0.72)

# =========================================================== 3b Probleme
s = slide("Probleme und Anpassungen", "Durchführung · Ist", "„Beim Umsetzen sind uns Abweichungen aufgefallen: [Problem 1 – was wich vom Plan ab?]. Wir haben deshalb [Anpassung 1] vorgenommen, weil [Begründung]. Außerdem: [Problem 2]. Dafür haben wir [Anpassung 2] gewählt, weil [Begründung].“")
for r, n in enumerate((1, 2)):
    y = 1.9 + r * 2.5
    for c, (h, b) in enumerate([("Problem", f"[Problem {n} – was wich vom Plan ab?]" if n == 1 else "[Problem 2]"), ("Anpassung", f"[Anpassung {n}]"), ("Begründung", "[Begründung]")]):
        x = 0.7 + c * 4.1
        card(s, x, y, 3.8, 2.1, h, b, 22, 20, fill=[BLUE_L, ACC_L, GREY_L][c], bar=[NAVY2_C, ACC_C, MUTED_C][c])

# =========================================================== 3c Rahmenbedingungen
s = slide("Rahmenbedingungen und Absprachen", "Durchführung", "„Unsere Rahmenbedingungen: ein Team aus zwei Personen, feste Abgabe am 8. Oktober, keine nennenswerten Kosten und Vertragsauszüge nur anonymisiert. Absprachen haben wir mit [Ansprechpartner] getroffen: [Thema und Ergebnis]. Berücksichtigt haben wir das, indem wir [Änderung] vorgenommen haben.“")
text(s, 0.7, 1.7, 5.8, 0.5, "Rahmenbedingungen", 24, True, NAVY_C)
for i, t in enumerate(["Team aus zwei Personen", "Feste Abgabe am 8. Oktober", "Keine nennenswerten Kosten", "Vertragsauszüge nur anonymisiert"]):
    rect(s, 0.7, 2.4 + i * 1.0, 5.8, 0.85, BLUE_L); rect(s, 0.7, 2.4 + i * 1.0, 0.1, 0.85, NAVY2_C)
    text(s, 1.0, 2.4 + i * 1.0, 5.4, 0.85, t, 22, True, NAVY_C, anchor=MSO_ANCHOR.MIDDLE)
text(s, 7.0, 1.7, 5.8, 0.5, "Absprachen", 24, True, NAVY_C)
rect(s, 7.0, 2.4, 5.8, 3.85, ACC_L); rect(s, 7.0, 2.4, 0.1, 3.85, ACC_C)
text(s, 7.3, 2.5, 5.4, 3.7, [[("Mit ", False, INK_C), ("[Ansprechpartner]", True, NAVY_C)], [("Thema und Ergebnis: ", False, INK_C), ("[Thema und Ergebnis]", True, NAVY_C)], [("Berücksichtigt durch: ", False, INK_C), ("[Änderung]", True, NAVY_C)]], 22)

# =========================================================== 3d Kundenorientierung
s = slide("Für wen wir das Booklet gebaut haben", "Kundenorientierung", "„Wir haben das Booklet für euch gebaut: in der Reihenfolge des Prüfungskatalogs, in einfacher Sprache statt Gesetzesdeutsch, mit Merksätzen, Aufgaben samt Lösungen und einem Merkblatt zum Mitnehmen. Aktuelle Zahlen wie die Mindestausbildungsvergütung 2026 sind mit Quellen belegt.“")
for i, (h, b) in enumerate([("Reihenfolge des Prüfungskatalogs", "Das Booklet folgt dem Prüfungskatalog."), ("Einfache Sprache", "statt Gesetzesdeutsch, mit Merksätzen"),
                            ("Aufgaben samt Lösungen", "und ein Merkblatt zum Mitnehmen"), ("Aktuelle Zahlen mit Quellen", "zum Beispiel die Mindestausbildungsvergütung 2026")]):
    card(s, 0.7 + (i % 2) * 6.15, 1.9 + (i // 2) * 2.4, 5.95, 2.1, h, b, 22, 20)

# =========================================================== 4 Booklet
s = slide("Unser Booklet: fünf Bausteine", "Das Booklet im Überblick", "„Unser Booklet hat fünf Bausteine. Erstens die Kapitel: zehn kurze Kapitel entlang des Prüfungskatalogs, von Beginn und Dauer über Arbeitszeit, Vergütung, Urlaub, Probezeit und Kündigung bis zu den Pflichten von Azubi und Ausbildenden. Zweitens Aufgaben im Prüfungsformat: eine Zuordnungsübung und fünf IHK-Fälle. Drittens der Erwartungshorizont: Zu jeder Aufgabe stehen Lösung, Paragraf und Begründung. Viertens Quizfragen zum Selbsttesten. Und fünftens das Merkblatt mit dem Wichtigsten auf einer Seite. Dazu kommen ein Glossar und ein Quellenverzeichnis.“")
s.shapes.add_picture(os.path.join(IMG, "b1.png"), Inches(9.5), Inches(1.5), height=Inches(5.3))
for i, (h, b) in enumerate([("Kapitel", "zehn kurze Kapitel entlang des Prüfungskatalogs"), ("Aufgaben", "Zuordnungsübung und fünf IHK-Fälle im Prüfungsformat"),
                            ("Erwartungshorizont", "Lösung, Paragraf und Begründung"), ("Quizfragen", "zum Selbsttesten"), ("Merkblatt", "das Wichtigste auf einer Seite")]):
    y = 1.6 + i * 1.05
    rect(s, 0.7, y, 0.8, 0.85, NAVY_C); text(s, 0.7, y, 0.8, 0.85, str(i + 1), 28, True, ACC_C, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    rect(s, 1.5, y, 7.6, 0.85, BLUE_L)
    text(s, 1.65, y, 7.4, 0.85, [[(h + ": ", True, NAVY_C), (b, False, INK_C)]], 19, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 6.9, 8.4, 0.4, "Dazu: Glossar und Quellenverzeichnis", 16, False, MUTED_C)

s = slide("Zwei Beispielseiten", "Das Booklet im Überblick", "„Zwei Beispiele zeigen wir euch kurz: das Kapitel zur Kündigung und einen IHK-Fall mit seinem Erwartungshorizont. Genau mit solchen Aufgaben trainiert ihr jetzt selbst.“")
for i, (f, cap) in enumerate([("b11.png", "Kapitel Kündigung"), ("b19.png", "IHK-Fälle mit Erwartungshorizont")]):
    x = 1.9 + i * 5.2
    s.shapes.add_picture(os.path.join(IMG, f), Inches(x), Inches(1.5), height=Inches(5.25))
    text(s, x - 0.3, 6.8, 4.2, 0.4, cap, 16, True, NAVY_C, PP_ALIGN.CENTER)

# =========================================================== Praxisteil Uebersicht
s = slide("Prüfungstraining mit der Klasse", "Praxisteil", "Der Prüfungsbereich Wirtschafts- und Sozialkunde dauert 60 Minuten, besteht aus fallbezogenen Aufgaben und zählt 10 % der Gesamtnote; laut IHK-Information sind es gebundene Aufgaben. Unsere Aufgaben ahmen das nach: kurzer Fall, Antworten A–D, genau eine Lösung.")
table(s, [["Station", "Format", "Dauer"],
          ["1  Wissens-Check", "Antwortkarten A–D, drei Blitzfragen", "3 Min."],
          ["2  Situationen zuordnen", "Kartenspiel in Zweierteams", "7 Min."],
          ["3  Wer wird Azubi-Millionär?", "Quizspiel, zwei Teams, Joker", "7 Min."],
          ["4  IHK-Fälle", "Gruppenarbeit mit Fallaufgaben", "7 Min."],
          ["5  Abschluss-Quiz", "Anonymes Exit-Ticket, fünf neue Situationen", "3 Min."]],
      0.7, 1.8, 12.0, [4.7, 5.6, 1.7], 22, 0.8)
text(s, 0.7, 6.85, 12, 0.5, "Wie in der Prüfung: kurzer Fall, Antworten A–D, genau eine Lösung · zusammen 27 Minuten", 18, False, MUTED_C)


# =========================================================== Lösungsfolie-Helfer
def solslide(station, rows, colw, head, size=20, rowh=0.8, title="Lösungen"):
    s = slide(title, station)
    table(s, [head] + rows, 0.7, 1.7, 12.0, colw, size, rowh)
    return s


# =========================================================== Station 1
intro("Station 1", "Wissens-Check", ["Frage auf der Folie, 15 Sekunden Bedenkzeit", "Auf „drei“ zeigen alle gleichzeitig ihre Karte A–D", "Danach die Lösung mit einem Satz Begründung"],
      "„Bevor wir starten, ein kurzer Check: Ihr bekommt drei Fragen und antwortet mit euren Karten A bis D. Auf ‚drei‘ zeigen alle gleichzeitig.“")
for i, q in enumerate(WISSENS_CHECK):
    question("Station 1", f"Frage {i+1} von 3", q["q"], q["a"])
solslide("Station 1 · Wissens-Check",
         [[f"{i+1}", f'{"ABCD"[q["c"]]} – {q["a"][q["c"]]}', q["fb"]] for i, q in enumerate(WISSENS_CHECK)],
         [1.2, 4.3, 6.5], ["Nr.", "Richtige Antwort", "Paragraf und Begründung"], 22, 0.95)

# =========================================================== Station 2
intro("Station 2", "Situationen zuordnen", ["Ihr seid die Ausbildungsberatung", "Legt acht Situationskarten zu acht Regelungskarten – ihr habt vier Minuten", "Danach Auflösung am Beamer; jedes Team zählt seine Treffer, die meisten gewinnen"],
      "„Ihr seid jetzt die Ausbildungsberatung: Ordnet jeder Situation die passende Regelung zu. Ihr habt vier Minuten.“")
s = slide("Lösungen", "Station 2 · Situationen zuordnen")
rows = [["Nr.", "Situation", "Passende Regelung"]]
for n, t, L in PAARE: rows.append([str(n), t, REGELUNGEN[L]])
table(s, rows, 0.7, 1.55, 12.0, [0.7, 5.5, 5.8], 14, 0.58)

# =========================================================== Station 3
intro("Station 3", "Wer wird Azubi-Millionär?", ["Zwei Teams, sieben Fragen im Wechsel", "Jede richtige Antwort bringt die nächste Gewinnstufe", "Wer falsch liegt, gibt die Frage an das andere Team weiter – es darf stehlen", "Jedes Team hat einmal den 50:50-Joker"],
      "„Willkommen bei ‚Wer wird Azubi-Millionär?‘ Zwei Teams, sieben Fragen, ein Joker pro Team. Wer falsch liegt, gibt die Frage an das andere Team weiter.“")
for i, q in enumerate(MILLIONAER):
    question("Azubi-Millionär", f"Frage {i+1} von 7", q["q"], q["a"], extra=q["stufe"])
solslide("Station 3 · Azubi-Millionär",
         [[m["stufe"], f'{"ABCD"[m["c"]]} – {m["a"][m["c"]]}', m["fb"]] for m in MILLIONAER],
         [1.8, 5.2, 5.0], ["Stufe", "Richtige Antwort", "Paragraf"], 17, 0.72)

# =========================================================== Station 4
intro("Station 4", "IHK-Fälle", ["Drei Gruppen bearbeiten je einen Fall wie in der Prüfung", "Jede Gruppe nennt Lösung und Begründung in einem Satz", "Schnelle Gruppen lösen einen der beiden Reservefälle"],
      "„Jetzt wird es prüfungsnah: Jede Gruppe bekommt einen Fall, wählt die richtige Antwort und begründet sie mit der passenden Regelung.“")
for f in FAELLE:
    question("Station 4", f'{"Reservefall" if f["reserve"] else "Fall"} {f["id"]} · {f["titel"]}', f["q"], f["a"])
solslide("Station 4 · IHK-Fälle",
         [[f["id"], "ABCD"[f["c"]], f["loesung"]] for f in FAELLE],
         [1.0, 1.3, 9.7], ["Fall", "Lösung", "Paragraf und Begründung"], 17, 0.95)

# =========================================================== Station 5
intro("Station 5", "Abschluss-Quiz", ["Fünf neue Situationen auf Folien", "Alle notieren anonym A–D auf einem Zettel", "Die Zettel werden eingesammelt und ausgezählt"],
      "„Zum Abschluss ein anonymes Mini-Quiz mit fünf neuen Situationen. Das Ergebnis zeigt uns, ob unser Booklet sein Ziel erreicht hat.“")
for i, q in enumerate(ABSCHLUSS):
    question("Station 5", f"Situation {i+1} von 5", q["q"], q["a"])
solslide("Station 5 · Abschluss-Quiz",
         [[str(i + 1), f'{"ABCD"[q["c"]]} – {q["a"][q["c"]]}', q["fb"]] for i, q in enumerate(ABSCHLUSS)],
         [0.9, 7.6, 3.5], ["Nr.", "Richtige Antwort", "Paragraf"], 18, 0.95)

# =========================================================== 5 Merkblatt
s = slide("Merkblatt: das Wichtigste auf einer Seite", "Zusammenfassung", "„Zum Mitnehmen bekommt ihr unser Merkblatt: das Wichtigste auf einer Seite. Die Kernregeln: Die Probezeit dauert ein bis vier Monate. Die Mindestausbildungsvergütung 2026 liegt bei 724, 854, 977 und 1.014 Euro. Mindesturlaub gibt es ab 18 Jahren mit 24 Werktagen, darunter mit 25, 27 oder 30. Gekündigt werden kann in der Probezeit jederzeit ohne Frist, danach fristlos aus wichtigem Grund oder durch den Azubi mit vier Wochen Frist beim Berufswechsel – immer schriftlich. Jugendliche arbeiten höchstens 8 Stunden täglich. Und der Vertragsinhalt wird seit August 2024 in Textform abgefasst.“")
rules = [("Probezeit", "ein bis vier Monate"), ("Vergütung 2026", "724 / 854 / 977 / 1.014 €"), ("Urlaub", "ab 18 Jahren 24 Werktage, darunter 25, 27 oder 30"),
         ("Kündigung", "Probezeit: jederzeit ohne Frist. Danach fristlos aus wichtigem Grund oder durch den Azubi mit 4 Wochen Frist beim Berufswechsel – immer schriftlich"),
         ("Jugendliche", "höchstens 8 Stunden täglich"), ("Vertrag", "seit August 2024 in Textform")]
for i, (h, b) in enumerate(rules):
    card(s, 0.7 + (i % 3) * 4.1, 1.7 + (i // 3) * 2.8, 3.85, 2.6, h, b, 22, 17 if i == 3 else 20, fill=ACC_L if i % 2 == 0 else BLUE_L, bar=ACC_C if i % 2 == 0 else NAVY2_C)

# =========================================================== 6 Evaluation
s = slide("Evaluation: Ziel erreicht?", "Ziel und Ergebnis", "„Hat das Projekt sein Ziel erreicht? Das Booklet deckt alle sechs Punkte des Prüfungskatalogs ab – Beginn und Dauer, tägliche Ausbildungszeit, Vergütung, Urlaub, Probezeit und Kündigung – jeweils mit Aufgaben und Lösung. Im Abschluss-Quiz haben die Mitschüler im Schnitt [Anzahl] von 5 Fragen richtig beantwortet. Die Zeit haben wir [eingehalten / nicht eingehalten], Kosten sind keine entstanden.\n\nDaraus leiten wir Verbesserungen ab: Erstens sollte das Booklet jährlich aktualisiert werden, denn die Mindestausbildungsvergütung wird jedes Jahr neu festgelegt. Zweitens planen wir das Probelesen durch Mitschüler beim nächsten Mal früher ein, damit Rückmeldungen noch ins Layout einfließen. Drittens: [eigener Vorschlag].“")
table(s, [["Ziel", "Ergebnis"],
          ["Alle sechs Punkte des Prüfungskatalogs", "Booklet deckt alle sechs ab – jeweils mit Aufgaben und Lösung"],
          ["Quiz-Durchschnitt (Abschluss-Quiz)", "[Anzahl] von 5 Fragen richtig"],
          ["Zeit", "[eingehalten / nicht eingehalten]"],
          ["Kosten", "keine entstanden"]], 0.7, 1.5, 12.0, [5.0, 7.0], 18, 0.52)
text(s, 0.7, 4.3, 12, 0.4, "Drei Verbesserungen", 20, True, NAVY_C)
for i, t in enumerate(["Booklet jährlich aktualisieren – die Mindestausbildungsvergütung wird jedes Jahr neu festgelegt", "Probelesen durch Mitschüler beim nächsten Mal früher einplanen", "[eigener Vorschlag]"]):
    card(s, 0.7 + i * 4.1, 4.8, 3.85, 2.4, f"{i+1}", t, 20, 17, fill=ACC_L, bar=ACC_C)

# =========================================================== 7 Reflexion
s = slide("Reflexion", "Gemeinsam und einzeln", "„Gemeinsam haben wir festgehalten: Gut lief [..]. Schwierig war [..]. Beim nächsten Projekt würden wir [..] anders machen.“\n\nEinzelreflexionen, je 0:30 (erst Yasin, dann Mido): „Fachlich habe ich gelernt, [..]. Methodisch nehme ich mit, [..].“")
for i, (h, b) in enumerate([("Was lief gut?", "[..]"), ("Was war schwierig?", "[..]"), ("Was machen wir beim nächsten Mal anders?", "[..]")]):
    card(s, 0.7 + i * 4.1, 2.0, 3.85, 3.6, h, b, 22, 28, off=1.5, fill=[BLUE_L, GREY_L, ACC_L][i], bar=[NAVY2_C, MUTED_C, ACC_C][i])

# =========================================================== 8 Fragen
s = slide(dark=True, notes="„Vielen Dank fürs Mitmachen! Zwei Fragen an euch: Was war am Booklet am hilfreichsten, und was fehlt euch noch? Und habt ihr Fragen zum Ausbildungsvertrag oder zu unserem Projekt?“")
text(s, 1.0, 1.0, 11, 0.5, "FRAGEN & KLASSENFEEDBACK", 16, True, ACC_C)
text(s, 1.0, 1.7, 11.3, 1.0, "Was war am Booklet am hilfreichsten, und was fehlt euch noch?", 34, True, WHITE)
text(s, 1.0, 3.5, 11.3, 1.0, "Habt ihr Fragen zum Ausbildungsvertrag oder zu unserem Projekt?", 34, True, WHITE)
rect(s, 1.0, 5.4, 2.0, 0.08, ACC_C)
text(s, 1.0, 5.6, 11, 1.0, "Vielen Dank!", 44, True, ACC_C)

out = os.path.join(os.path.dirname(__file__), "..", "output", "praesentation.pptx")
prs.save(out)
print("Folien:", len(prs.slides))
