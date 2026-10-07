# -*- coding: utf-8 -*-
import os, sys, subprocess
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
OUT = os.path.join(os.path.dirname(__file__), "..", "output", "v4")
NAVY, NAVY2, ACC = "14264B", "22386B", "F5A800"
FONT = "Arial"

def rgb(h): return RGBColor.from_string(h)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); tcPr.append(sh)

def cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        v = kw.get(edge)
        e = OxmlElement(f"w:{edge}")
        if v: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v[0])); e.set(qn("w:color"), v[1])
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def fix_widths(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for c, w in zip(row.cells, widths): c.width = Cm(w)

def run(p, text, size=10, bold=False, color="1A2238", italic=False):
    r = p.add_run(text); r.font.name = FONT; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = rgb(color); return r

def para(container, text="", size=10, bold=False, color="1A2238", after=3, before=0, align=None, line=1.15):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before); p.paragraph_format.line_spacing = line
    if align: p.alignment = align
    if text: run(p, text, size, bold, color)
    return p

def heading(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); l = OxmlElement("w:left")
    l.set(qn("w:val"), "single"); l.set(qn("w:sz"), "36"); l.set(qn("w:space"), "6"); l.set(qn("w:color"), ACC); bd.append(l); pPr.append(bd)
    p.paragraph_format.left_indent = Cm(0.25)
    run(p, f"{num}  ", 11.5, True, ACC.replace("F5A800", "B07A00")); run(p, text, 11.5, True, NAVY)
    return p

def bullet(doc, parts, size=10, after=1.5, indent=0.55):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent); p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.12
    run(p, "•  ", size, True, "B07A00")
    for t, b in parts: run(p, t, size, b)
    return p

def table(doc, widths, header, rows, size=9.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(widths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.width = Cm(widths[j]); shade(c, NAVY); cell_margins(c, 50, 50)
        c.paragraphs[0].paragraph_format.space_after = Pt(0); run(c.paragraphs[0], h, size, True, "FFFFFF")
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.rows[i + 1].cells[j]; c.width = Cm(widths[j]); cell_margins(c, 45, 45)
            shade(c, "F6F8FC" if i % 2 == 0 else "FFFFFF"); cell_borders(c, bottom=(4, "D5D9E3"))
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            run(p, v, size, j == 0, NAVY if j == 0 else "1A2238")
    fix_widths(t, widths)
    return t

def goal_box(doc, title, text, examples):
    t = doc.add_table(rows=1, cols=2); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.rows[0].cells
    a.width, b.width = Cm(0.35), Cm(16.65)
    shade(a, ACC); shade(b, "FFF8E1"); cell_margins(b, 90, 90, 160, 140)
    p = b.paragraphs[0]; p.paragraph_format.space_after = Pt(2); run(p, title, 11, True, NAVY)
    p2 = b.add_paragraph(); p2.paragraph_format.space_after = Pt(3); p2.paragraph_format.line_spacing = 1.15; run(p2, text, 10)
    p3 = b.add_paragraph(); p3.paragraph_format.space_after = Pt(1); run(p3, "Zum Beispiel:", 10, True, "B07A00")
    for s_, r_ in examples:
        q = b.add_paragraph(); q.paragraph_format.space_after = Pt(0.5); q.paragraph_format.left_indent = Cm(0.4); q.paragraph_format.line_spacing = 1.1
        run(q, s_, 9.5, True, NAVY); run(q, "  →  " + r_, 9.5)
    fix_widths(t, [0.35, 16.65])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def field(p, code, size=8, color="5B6680"):
    r = p.add_run(); r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = rgb(color)
    for tp, txt in (("begin", None), (None, code), ("end", None)):
        if tp:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tp); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)



def new_doc():
    d = Document(); sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.5)
    st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return d, sec

def title_block(doc, kicker, title, sub):
    tt = doc.add_table(rows=1, cols=2); tt.autofit = False; tt.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = tt.rows[0].cells; a.width, b.width = Cm(0.5), Cm(16.5)
    shade(a, ACC); shade(b, NAVY); cell_margins(b, 150, 150, 260, 200); fix_widths(tt, [0.5, 16.5])
    p = b.paragraphs[0]; p.paragraph_format.space_after = Pt(2); run(p, kicker, 9, True, ACC)
    p = b.add_paragraph(); p.paragraph_format.space_after = Pt(2); run(p, title, 19, True, "FFFFFF")
    p = b.add_paragraph(); p.paragraph_format.space_after = Pt(0); run(p, sub, 9.5, False, "DCE3F2")

def footer(sec, label):
    sec.footer.paragraphs[0].style.font.size = Pt(8); sec.footer.paragraphs[0].style.font.name = FONT
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, label + " · Seite ", 8, False, "5B6680"); field(fp, "PAGE"); run(fp, " von ", 8, False, "5B6680"); field(fp, "NUMPAGES")


def to_pdf(path, outdir=None):
    import pymupdf
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir or os.path.dirname(path), path], capture_output=True, timeout=240)
    return len(pymupdf.open(path.replace(".docx", ".pdf")))
