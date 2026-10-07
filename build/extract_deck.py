# -*- coding: utf-8 -*-
import os, re, json
from lxml import etree
from pptx import Presentation
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
X = lambda el, e: etree._Element.xpath(el, e, namespaces=NS)
def load(path=None):
    P = Presentation(path or os.path.join(ROOT, "output/v3/praesentation_unterricht_v4.pptx")); out = []
    for i, s in enumerate(P.slides, 1):
        note = s.notes_slide.notes_text_frame.text
        g = lambda k: (re.search(k + r": (.*)", note) or [None, ""])[1].strip()
        names = {sh.shape_id: sh.name for sh in s.shapes}; clicks, auto = [], []
        t = s._element.find(".//p:timing", NS)
        if t is not None:
            for outer in X(t, ".//p:cTn[@nodeType='mainSeq']/p:childTnLst/p:par"):
                is_auto = any(c.get("evt") == "onBegin" for c in X(outer, "./p:cTn/p:stCondLst/p:cond"))
                grp = [names[int(X(e, ".//p:spTgt/@spid")[0])] for e in X(outer, ".//p:cTn[@presetClass='entr']")]
                (auto.extend(grp) if is_auto else clicks.append(grp))
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        d = g("DAUER"); dur = int(d.split(":")[0]) * 60 + int(d.split(":")[1]) if re.match(r"\d+:\d+", d) else None
        phase = re.search(r"PHASE: (.*?) · FOLIENTYP: (.*)", note)
        out.append(dict(nr=i, hidden=s._element.get("show") == "0", phase=phase.group(1), kind=phase.group(2), spr=g("SPRECHER"), dauer=dur, sozial=g("SOZIALFORM"),
                        stich=g("IMPULS / MODERATION"), erw=g("ERWARTETE ANTWORTEN"), puffer=g("ZEITPUFFER / STREICHEN"), clicks=clicks, auto=auto, texts=texts, notes=note))
    return out
if __name__ == "__main__":
    D = load()
    for d in D: print(d["nr"], "H" if d["hidden"] else " ", d["phase"], "|", d["kind"], "|", d["spr"], d["dauer"], "| clicks", len(d["clicks"]), "auto", len(d["auto"]), "|", d["texts"][6][:50].replace("\n", " ") if len(d["texts"]) > 6 else d["texts"][-1][:50])
