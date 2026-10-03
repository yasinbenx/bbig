# -*- coding: utf-8 -*-
import os, sys, subprocess
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import plan2
from plan2 import *

OUT = os.path.join(os.path.dirname(__file__), "..", "output", "v2")
NAVY, NAVY2, ACC = "14264B", "22386B", "F5A800"
FONT = "Arial"

def rgb(h): return RGBColor.from_string(h)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); tcPr.append(sh)

def cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        v = kw.get(edge)
        e = OxmlElement(f"w:{edge}")
        if v: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v[0])); e.set(qn("w:color"), v[1])
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def fix_widths(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for c, w in zip(row.cells, widths): c.width = Cm(w)

def run(p, text, size=10, bold=False, color="1A2238", italic=False):
    r = p.add_run(text); r.font.name = FONT; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = rgb(color); return r

def para(container, text="", size=10, bold=False, color="1A2238", after=3, before=0, align=None, line=1.15):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before); p.paragraph_format.line_spacing = line
    if align: p.alignment = align
    if text: run(p, text, size, bold, color)
    return p

def heading(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); l = OxmlElement("w:left")
    l.set(qn("w:val"), "single"); l.set(qn("w:sz"), "36"); l.set(qn("w:space"), "6"); l.set(qn("w:color"), ACC); bd.append(l); pPr.append(bd)
    p.paragraph_format.left_indent = Cm(0.25)
    run(p, f"{num}  ", 11.5, True, ACC.replace("F5A800", "B07A00")); run(p, text, 11.5, True, NAVY)
    return p

def bullet(doc, parts, size=10, after=1.5, indent=0.55):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent); p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.12
    run(p, "•  ", size, True, "B07A00")
    for t, b in parts: run(p, t, size, b)
    return p

def table(doc, widths, header, rows, size=9.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(widths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.width = Cm(widths[j]); shade(c, NAVY); cell_margins(c, 50, 50)
        c.paragraphs[0].paragraph_format.space_after = Pt(0); run(c.paragraphs[0], h, size, True, "FFFFFF")
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.rows[i + 1].cells[j]; c.width = Cm(widths[j]); cell_margins(c, 45, 45)
            shade(c, "F6F8FC" if i % 2 == 0 else "FFFFFF"); cell_borders(c, bottom=(4, "D5D9E3"))
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            run(p, v, size, j == 0, NAVY if j == 0 else "1A2238")
    fix_widths(t, widths)
    return t

def goal_box(doc, title, text, examples):
    t = doc.add_table(rows=1, cols=2); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.rows[0].cells
    a.width, b.width = Cm(0.35), Cm(16.65)
    shade(a, ACC); shade(b, "FFF8E1"); cell_margins(b, 90, 90, 160, 140)
    p = b.paragraphs[0]; p.paragraph_format.space_after = Pt(2); run(p, title, 11, True, NAVY)
    p2 = b.add_paragraph(); p2.paragraph_format.space_after = Pt(3); p2.paragraph_format.line_spacing = 1.15; run(p2, text, 10)
    p3 = b.add_paragraph(); p3.paragraph_format.space_after = Pt(1); run(p3, "Zum Beispiel:", 10, True, "B07A00")
    for s_, r_ in examples:
        q = b.add_paragraph(); q.paragraph_format.space_after = Pt(0.5); q.paragraph_format.left_indent = Cm(0.4); q.paragraph_format.line_spacing = 1.1
        run(q, s_, 9.5, True, NAVY); run(q, "  →  " + r_, 9.5)
    fix_widths(t, [0.35, 16.65])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def field(p, code, size=8, color="5B6680"):
    r = p.add_run(); r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = rgb(color)
    for tp, txt in (("begin", None), (None, code), ("end", None)):
        if tp:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tp); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)


import pymupdf
from data2 import FEHLER, LEHRKRAFT
AP_N = {n: nm for n, nm, *_ in AP}

def new_doc():
    d = Document(); sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.5)
    st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return d, sec

def title_block(doc, kicker, title, sub):
    tt = doc.add_table(rows=1, cols=2); tt.autofit = False; tt.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = tt.rows[0].cells; a.width, b.width = Cm(0.5), Cm(16.5)
    shade(a, ACC); shade(b, NAVY); cell_margins(b, 150, 150, 260, 200); fix_widths(tt, [0.5, 16.5])
    p = b.paragraphs[0]; p.paragraph_format.space_after = Pt(2); run(p, kicker, 9, True, ACC)
    p = b.add_paragraph(); p.paragraph_format.space_after = Pt(2); run(p, title, 19, True, "FFFFFF")
    p = b.add_paragraph(); p.paragraph_format.space_after = Pt(0); run(p, sub, 9.5, False, "DCE3F2")

def footer(sec, label):
    sec.footer.paragraphs[0].style.font.size = Pt(8); sec.footer.paragraphs[0].style.font.name = FONT
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, label + " · Seite ", 8, False, "5B6680"); field(fp, "PAGE"); run(fp, " von ", 8, False, "5B6680"); field(fp, "NUMPAGES")

def to_pdf(path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, path], capture_output=True, timeout=240)
    return len(pymupdf.open(path.replace(".docx", ".pdf")))

# ============================================================ Projektskizze
doc, sec = new_doc()
title_block(doc, "PROJEKTSKIZZE", "Rechte und Pflichten aus dem Ausbildungsvertrag", "Fach GP (Geschäftsprozesse) · Kaufleute für Büromanagement · Präsentation am 08.10.2026")
heading(doc, "1.", "Vorläufiger Arbeitstitel")
para(doc, "Rechte und Pflichten aus dem Ausbildungsvertrag – Booklet und interaktive Stunde zur Prüfungsvorbereitung im Fach GP", 10)
heading(doc, "2.", "Einreicher / Ansprechpartner")
p = para(doc, "", 10); run(p, "Yasin & Mido", 10, True, NAVY); run(p, f"  ·  Betreuende Lehrkraft: {LEHRKRAFT}  ·  Klasse: [eintragen]", 10)
heading(doc, "3.", "Vorläufige Umschreibung der Projektidee")
para(doc, "Wir werden ein Booklet und eine interaktive Unterrichtsstunde zum Thema „Rechte und Pflichten aus dem Ausbildungsvertrag“ erstellen, um Mitschüler auf die schriftliche IHK-Abschlussprüfung am 25.11.2026 vorzubereiten (Prüfungsbereich Wirtschafts- und Sozialkunde). Inhaltlich geht es um die sechs Prüfungspunkte Beginn und Dauer der Ausbildung, tägliche Ausbildungszeit, Vergütung, Urlaub, Probezeit und Kündigung sowie um die Pflichten von Auszubildenden und Ausbildenden. Rechtsgrundlage ist vor allem das Berufsbildungsgesetz (BBiG).", 10, after=3)
para(doc, "Das Booklet soll 13 Seiten umfassen: kurze Kapitel in einfacher Sprache, ein Merkblatt für eine Seite, Aufgaben (Vertrags-Detektiv, IHK-Fälle, Wahr oder falsch) mit Erwartungshorizont, ein Mini-Quiz, Glossar und Quellenverzeichnis. Die Praxisbeispiele sollen aus einem anonymisierten, eigenen Ausbildungsvertrag stammen.", 10, after=3)
para(doc, "In der Unterrichtsstunde (bis zu 45 Minuten) stellen wir das Projekt kurz vor und üben danach mit der Klasse in zwei Aktivitäten: Rote/Grüne Karte und Vertrags-Detektiv. Dazu schreiben wir einen Projektbericht.", 10, after=2)
heading(doc, "4.", "Vorteile / erwarteter Nutzen")
bullet(doc, [("Mitschüler: ", True), ("Zum Ausbildungsvertrag gibt es bisher kein kompaktes Lernmaterial mit Aufgaben und Lösungen. Unser Booklet erklärt die Regeln in einfacher Sprache statt Gesetzesdeutsch.", False)])
bullet(doc, [("Prüfungsnah: ", True), ("Der IHK-Prüfungsbereich Wirtschafts- und Sozialkunde dauert 60 Minuten, besteht aus fallbezogenen Aufgaben und zählt 10 % der Gesamtnote. Unsere Aufgaben üben genau solche Fälle.", False)])
bullet(doc, [("Interaktiv: ", True), ("Die Klasse lernt durch Mitmachen statt durch Zuhören.", False)])
bullet(doc, [("Wir selbst: ", True), ("Wir lernen die Regeln aus dem eigenen Ausbildungsvertrag sicher anzuwenden und setzen die Methoden aus dem Projektmanagement-Unterricht praktisch um.", False)], after=2)
heading(doc, "5.", "Umsetzung")
para(doc, "Die Arbeit gliedert sich in vier Stufen, die im Projektstrukturplan als Teilaufgaben und Arbeitspakete (AP) wiederkehren:", 10, after=3)
table(doc, [0.8, 4.9, 9.6, 1.7], ["", "Stufe", "Inhalt", "AP"], [
    ("1", "Vorbereitung", "Projektskizze und Ziele erstellen; Planung mit Projektstrukturplan und Projektablaufplan", "1–3"),
    ("2", "Praxisbeispiele", "Praxisbeispiele sammeln und anonymisieren", "4"),
    ("3", "Booklet und Aufgaben", "Booklet mit Kapiteln, Merkblatt und Mini-Quiz erstellen; Aufgaben mit Erwartungshorizont erstellen", "5–6"),
    ("4", "Präsentation und Abschluss", "Präsentation erstellen; Material, Korrekturlesen und Probelauf; Projektbericht schreiben; Präsentation halten (08.10.)", "7–10"),
])
para(doc, "Aufgabenverteilung: Recherche und Texte liegen bei Yasin; Vertragssichtung, Praxisbeispiele, Grafiken und Korrekturlesen bei Mido; Gliederung, Layout und Präsentation machen wir gemeinsam. Projektstrukturplan und Projektablaufplan liegen im Anhang.", 9.5, after=2, before=4)
heading(doc, "6.", "Laufzeit")
para(doc, "4 Wochen (10.09. – 08.10.2026), geplant in Arbeitstagen Montag bis Freitag; die Umsetzung findet im Oktober statt:", 10, after=3)
table(doc, [4.7, 10.5, 1.8], ["Zeitraum", "Schwerpunkt", "AP"], [
    ("Woche 1 (10.09.–16.09.)", "Projektskizze und Ziele erstellen", "1"),
    ("Woche 2 (17.09.–23.09.)", "Projektstrukturplan erstellen, Projektablaufplan beginnen, Praxisbeispiele sammeln", "2–4"),
    ("Woche 3 (24.09.–30.09.)", "Projektablaufplan abschließen, Praxisbeispiele sammeln und anonymisieren", "3–4"),
    ("Woche 4 (01.10.–07.10.)", "Umsetzung: Booklet, Aufgaben, Präsentation, Material und Probelauf, Projektbericht", "5–9"),
    ("08.10.2026", "Präsentation halten", "10"),
])
# Kontrolle: Wochenzuordnung stimmt mit dem berechneten Plan ueberein
import datetime as _dt
def _aps(a, b): return sorted(n for n in ES if start(n) <= b and end(n) >= a)
assert _aps(_dt.date(2026,9,10), _dt.date(2026,9,16)) == [1], _aps(_dt.date(2026,9,10), _dt.date(2026,9,16))
assert _aps(_dt.date(2026,9,17), _dt.date(2026,9,23)) == [2,3,4], _aps(_dt.date(2026,9,17), _dt.date(2026,9,23))
assert _aps(_dt.date(2026,9,24), _dt.date(2026,9,30)) == [3,4], _aps(_dt.date(2026,9,24), _dt.date(2026,9,30))
assert _aps(_dt.date(2026,10,1), _dt.date(2026,10,7)) == [5,6,7,8,9], _aps(_dt.date(2026,10,1), _dt.date(2026,10,7))
ms = " · ".join(f"M{k} {nm} ({fmt(t)})" for k, nm, t in MS)
para(doc, "Meilensteine: " + ms + ".", 9.5, after=2, before=4)
heading(doc, "7.", "Finanzierung")
para(doc, "Für das Projekt entstehen keine nennenswerten Kosten. Benötigt werden digitale Recherchemittel, ein Computer und ein Canva-Zugang für das Layout (vorhanden). Für den Druck von Booklet, Merkblatt und Material: [Anzahl der Ausdrucke und Drucker klären].", 10, after=2)
h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(4); h.paragraph_format.keep_with_next = True
run(h, "Unsere zwei SMART-Ziele", 13, True, NAVY)
goal_box(doc, "Ziel 1 – Wissen über den Ausbildungsvertrag",
         "Bis zum 08.10.2026 kennen wir die wichtigsten Regelungen aus Berufsbildungsgesetz und Ausbildungsvertrag und können für jeden der sechs Prüfungspunkte (Beginn und Dauer, tägliche Ausbildungszeit, Vergütung, Urlaub, Probezeit, Kündigung) mindestens ein Beispiel erklären.",
         [("Beginn und Dauer", "Verkürzung auf gemeinsamen Antrag von Azubi und Ausbildenden (§ 8 Abs. 1 BBiG)"),
          ("Tägliche Ausbildungszeit", "Jugendliche höchstens 8 Stunden täglich (§ 8 JArbSchG)"),
          ("Vergütung", "mindestens 724 € im 1. Jahr bei Ausbildungsbeginn 2026 (§ 17 BBiG)"),
          ("Urlaub", "ab 18 Jahren mindestens 24 Werktage, darunter 25, 27 oder 30 (§ 3 BUrlG, § 19 JArbSchG)"),
          ("Probezeit", "ein bis vier Monate (§ 20 BBiG)"),
          ("Kündigung", "in der Probezeit jederzeit ohne Frist, aber schriftlich (§ 22 Abs. 1 und 3 BBiG)")])
goal_box(doc, "Ziel 2 – Anwendung im Ausbildungsalltag",
         "Bis zur Präsentation am 08.10.2026 finden Mitschüler im Vertrags-Detektiv mindestens 4 von 5 Fehlern und begründen sie mit der passenden Regel.",
         [("Probezeit von sechs Monaten", FEHLER[1][1]), ("650 € im 1. Ausbildungsjahr", FEHLER[2][1]),
          ("9 Stunden tägliche Ausbildungszeit", FEHLER[3][1]), ("20 Werktage Urlaub", FEHLER[4][1]),
          ("Kündigung in der Probezeit nur mit vier Wochen Frist", FEHLER[5][1])])
footer(sec, "Projektskizze · Yasin & Mido")
path = os.path.join(OUT, "projektskizze.docx"); doc.save(path); print("Skizze Seiten:", to_pdf(path))

# ============================================================ Protokoll-Vorlage
doc, sec = new_doc(); sec.top_margin = Cm(1.2); sec.bottom_margin = Cm(1.3)
title_block(doc, "BESPRECHUNG", "Protokoll", "Rechte und Pflichten aus dem Ausbildungsvertrag · Yasin & Mido")
def kv(rows, w1=4.2, w2=12.8, h=0.85):
    t = doc.add_table(rows=len(rows), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v, hh) in enumerate(rows):
        a, b = t.rows[i].cells; shade(a, "E8EEF9"); cell_margins(a, 70, 70); cell_margins(b, 70, 70)
        for c in (a, b): cell_borders(c, top=(4, "BFC6D6"), bottom=(4, "BFC6D6"), left=(4, "BFC6D6"), right=(4, "BFC6D6"))
        a.paragraphs[0].paragraph_format.space_after = Pt(0); run(a.paragraphs[0], k, 9.5, True, NAVY)
        b.paragraphs[0].paragraph_format.space_after = Pt(0)
        if v: run(b.paragraphs[0], v, 9.5)
        tr = t.rows[i]._tr; trPr = tr.get_or_add_trPr(); he = OxmlElement("w:trHeight"); he.set(qn("w:val"), str(int((hh or h) * 567))); he.set(qn("w:hRule"), "atLeast"); trPr.append(he)
    fix_widths(t, [w1, w2]); return t
doc.add_paragraph().paragraph_format.space_after = Pt(2)
kv([("Projekt", "Rechte und Pflichten aus dem Ausbildungsvertrag", None), ("Datum", "", None), ("Uhrzeit", "", None), ("Ort / Form", "", None), ("Teilnehmer", "", None), ("Thema", "", None)])
doc.add_paragraph().paragraph_format.space_after = Pt(2)
kv([("Besprochene Punkte", "", 3.6), ("Ergebnisse / Beschlüsse", "", 3.2)])
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(3); run(p, "Aufgaben", 11, True, NAVY)
t = doc.add_table(rows=5, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, hname in enumerate(["Wer", "Was", "Bis wann"]):
    c = t.rows[0].cells[j]; shade(c, NAVY); cell_margins(c, 50, 50); c.paragraphs[0].paragraph_format.space_after = Pt(0); run(c.paragraphs[0], hname, 9.5, True, "FFFFFF")
for i in range(1, 5):
    for c in t.rows[i].cells:
        cell_margins(c, 90, 90); cell_borders(c, bottom=(4, "BFC6D6"), left=(4, "BFC6D6"), right=(4, "BFC6D6")); c.paragraphs[0].paragraph_format.space_after = Pt(0)
fix_widths(t, [3.6, 10.2, 3.2])
doc.add_paragraph().paragraph_format.space_after = Pt(2)
kv([("Offene Punkte", "", 1.9), ("Nächster Termin", "", None), ("Protokollführer", "", None)])
footer(sec, "Protokoll · Yasin & Mido")
path = os.path.join(OUT, "protokoll_vorlage.docx"); doc.save(path); print("Protokoll Seiten:", to_pdf(path))
