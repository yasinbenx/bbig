# -*- coding: utf-8 -*-
import os, sys, subprocess
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from plan import *

OUT = os.path.join(os.path.dirname(__file__), "..", "output", "projektplanung")
NAVY, NAVY2, ACC = "14264B", "22386B", "F5A800"
FONT = "Arial"

def rgb(h): return RGBColor.from_string(h)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); tcPr.append(sh)

def cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        v = kw.get(edge)
        e = OxmlElement(f"w:{edge}")
        if v: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v[0])); e.set(qn("w:color"), v[1])
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def fix_widths(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for c, w in zip(row.cells, widths): c.width = Cm(w)

def run(p, text, size=10, bold=False, color="1A2238", italic=False):
    r = p.add_run(text); r.font.name = FONT; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = rgb(color); return r

def para(container, text="", size=10, bold=False, color="1A2238", after=3, before=0, align=None, line=1.15):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before); p.paragraph_format.line_spacing = line
    if align: p.alignment = align
    if text: run(p, text, size, bold, color)
    return p

def heading(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); l = OxmlElement("w:left")
    l.set(qn("w:val"), "single"); l.set(qn("w:sz"), "36"); l.set(qn("w:space"), "6"); l.set(qn("w:color"), ACC); bd.append(l); pPr.append(bd)
    p.paragraph_format.left_indent = Cm(0.25)
    run(p, f"{num}  ", 11.5, True, ACC.replace("F5A800", "B07A00")); run(p, text, 11.5, True, NAVY)
    return p

def bullet(doc, parts, size=10, after=1.5, indent=0.55):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent); p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.12
    run(p, "•  ", size, True, "B07A00")
    for t, b in parts: run(p, t, size, b)
    return p

def table(doc, widths, header, rows, size=9.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(widths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.width = Cm(widths[j]); shade(c, NAVY); cell_margins(c, 50, 50)
        c.paragraphs[0].paragraph_format.space_after = Pt(0); run(c.paragraphs[0], h, size, True, "FFFFFF")
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.rows[i + 1].cells[j]; c.width = Cm(widths[j]); cell_margins(c, 45, 45)
            shade(c, "F6F8FC" if i % 2 == 0 else "FFFFFF"); cell_borders(c, bottom=(4, "D5D9E3"))
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            run(p, v, size, j == 0, NAVY if j == 0 else "1A2238")
    fix_widths(t, widths)
    return t

def goal_box(doc, title, text, examples):
    t = doc.add_table(rows=1, cols=2); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.rows[0].cells
    a.width, b.width = Cm(0.35), Cm(16.65)
    shade(a, ACC); shade(b, "FFF8E1"); cell_margins(b, 90, 90, 160, 140)
    p = b.paragraphs[0]; p.paragraph_format.space_after = Pt(2); run(p, title, 11, True, NAVY)
    p2 = b.add_paragraph(); p2.paragraph_format.space_after = Pt(3); p2.paragraph_format.line_spacing = 1.15; run(p2, text, 10)
    p3 = b.add_paragraph(); p3.paragraph_format.space_after = Pt(1); run(p3, "Zum Beispiel:", 10, True, "B07A00")
    for s_, r_ in examples:
        q = b.add_paragraph(); q.paragraph_format.space_after = Pt(0.5); q.paragraph_format.left_indent = Cm(0.4); q.paragraph_format.line_spacing = 1.1
        run(q, s_, 9.5, True, NAVY); run(q, "  →  " + r_, 9.5)
    fix_widths(t, [0.35, 16.65])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def field(p, code, size=8, color="5B6680"):
    r = p.add_run(); r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = rgb(color)
    for tp, txt in (("begin", None), (None, code), ("end", None)):
        if tp:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tp); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.5)
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

# Titelblock
tt = doc.add_table(rows=1, cols=2); tt.autofit = False; tt.alignment = WD_TABLE_ALIGNMENT.CENTER
a, b = tt.rows[0].cells; a.width, b.width = Cm(0.5), Cm(16.5)
shade(a, ACC); shade(b, NAVY); cell_margins(b, 150, 150, 260, 200)
fix_widths(tt, [0.5, 16.5])
p = b.paragraphs[0]; p.paragraph_format.space_after = Pt(2); run(p, "PROJEKTSKIZZE", 9, True, ACC)
p = b.add_paragraph(); p.paragraph_format.space_after = Pt(2); run(p, "Rechte und Pflichten aus dem Ausbildungsvertrag", 19, True, "FFFFFF")
p = b.add_paragraph(); p.paragraph_format.space_after = Pt(0); run(p, "Themenschwerpunkt Verkürzer · Kaufleute für Büromanagement · Präsentation am 08.10.2026", 9.5, False, "DCE3F2")

# 1
heading(doc, "1.", "Vorläufiger Arbeitstitel")
para(doc, "Rechte und Pflichten aus dem Ausbildungsvertrag – Prüfungsvorbereitungsbooklet für die IHK-Abschlussprüfung in Wirtschafts- und Sozialkunde", 10)
# 2
heading(doc, "2.", "Einreicher / Ansprechpartner")
p = para(doc, "", 10); run(p, "Yasin & Mido", 10, True, NAVY); run(p, "  ·  [Klasse und Kontakt eintragen]", 10)
# 3
heading(doc, "3.", "Vorläufige Umschreibung der Projektidee")
para(doc, "Wir werden ein Prüfungsvorbereitungsbooklet zu einem Thema der schriftlichen IHK-Abschlussprüfung in Wirtschafts- und Sozialkunde erstellen: „Rechte und Pflichten aus dem Ausbildungsvertrag“ (Unterpunkt 0103 im Prüfungskatalog). Auf 25 Seiten soll es die sechs Prüfungsinhalte – Beginn und Dauer der Ausbildung, tägliche Ausbildungszeit, Vergütung, Urlaub, Probezeit und Kündigung – sowie die Pflichten von Auszubildenden und Ausbildenden in einfacher Sprache erklären. Rechtsgrundlage ist vor allem das Berufsbildungsgesetz (BBiG).", 10, after=3)
para(doc, "Das Booklet wird ein Merkblatt für eine Seite, Aufgaben im Prüfungsformat mit Erwartungshorizont (Lösung, Paragraf, Begründung), drei Quizrunden sowie Glossar und Quellenverzeichnis enthalten. Die Praxisbeispiele sollen aus einem anonymisierten, eigenen Ausbildungsvertrag stammen.", 10, after=3)
para(doc, "Ergänzend planen wir eine 45-minütige Präsentation mit anschließendem Prüfungstraining für die Klasse (fünf Stationen), ein Quiz als HTML-Datei und einen Projektbericht.", 10, after=2)
# 4
heading(doc, "4.", "Vorteile / erwarteter Nutzen")
bullet(doc, [("Mitschüler: ", True), ("Zum Ausbildungsvertrag gibt es bisher kein kompaktes Lernmaterial mit Aufgaben und Lösungen. Unser Booklet erklärt die Regeln in einfacher Sprache statt Gesetzesdeutsch und bereitet auf die Abschlussprüfung am 25.11.2026 vor.", False)])
bullet(doc, [("Prüfungsnah: ", True), ("Wirtschafts- und Sozialkunde dauert 60 Minuten, besteht aus fallbezogenen Aufgaben und zählt 10 % der Gesamtnote. Unsere Aufgaben ahmen das nach: kurzer Fall, Antworten A–D, genau eine Lösung.", False)])
bullet(doc, [("Wir selbst: ", True), ("Wir lernen die Regeln aus dem eigenen Ausbildungsvertrag sicher anzuwenden und wenden die Methoden aus dem Projektmanagement-Unterricht praktisch an.", False)], after=2)
# 5
heading(doc, "5.", "Umsetzung")
para(doc, "Wir arbeiten nach dem Wasserfall-Modell, also Phase für Phase, weil Anforderungen, Termin und Ergebnis von Anfang an feststehen. Die Arbeit gliedert sich in sechs Stufen, die im Projektstrukturplan als Teilaufgaben und Arbeitspakete (AP) wiederkehren:", 10, after=3)
table(doc, [0.8, 4.6, 9.9, 1.7], ["", "Stufe", "Inhalt", "AP"], [
    ("1", "Planung", "Problemanalyse und Zielsetzung, Projektplanung", "1–2"),
    ("2", "Recherche & Konzeption", "BBiG und Berufsausbildungsvertrag recherchieren, eigenen Vertrag sichten und anonymisieren, Gliederung erstellen", "3–5"),
    ("3", "Booklet", "Texte, Praxisbeispiele und Grafiken, Layout; Übungsteil mit Merkblatt, Aufgaben samt Erwartungshorizont und Quizfragen", "6–11"),
    ("4", "Präsentation & Praxisteil", "Präsentation mit Prüfungstraining, Quiz als HTML-Datei, Druckmaterial, Generalprobe", "12–15"),
    ("5", "Qualitätssicherung", "Fachlicher Abgleich mit dem BBiG, Korrekturlesen und Probelesen", "16–17"),
    ("6", "Projektbericht & Abgabe", "Projektbericht schreiben, Abgabe und Präsentation am 08.10.2026", "18–19"),
])
para(doc, "Aufgabenverteilung: Recherche und Texte liegen bei Yasin; Vertragssichtung, Praxisbeispiele, Grafiken und Korrekturlesen bei Mido; Gliederung, Layout und Präsentation machen wir gemeinsam. Projektstrukturplan und Projektablaufplan liegen im Anhang.", 9.5, after=2, before=4)
# 6
heading(doc, "6.", "Laufzeit")
para(doc, "4 Wochen (10.09. – 08.10.2026), geplant in Arbeitstagen Montag bis Freitag:", 10, after=3)
table(doc, [4.7, 10.5, 1.8], ["Zeitraum", "Schwerpunkt", "AP"], [
    ("Woche 1 (10.09.–16.09.)", "Planung, Recherche, Vertragssichtung, Beginn der Gliederung", "1–5"),
    ("Woche 2 (17.09.–23.09.)", "Gliederung abschließen, Texte und Praxisbeispiele beginnen, Projektbericht beginnen", "5–7, 18"),
    ("Woche 3 (24.09.–30.09.)", "Texte und Beispiele abschließen, Merkblatt, Aufgaben mit Erwartungshorizont, Quizfragen; Start von Layout, Präsentation, Quiz-Datei und Druckmaterial", "6–14, 18"),
    ("Woche 4 (01.10.–07.10.)", "Layout, Präsentation, Quiz-Datei, Druckmaterial, Generalprobe, Qualitätssicherung, Projektbericht", "8, 12–18"),
    ("08.10.2026", "Abgabe und Präsentation", "19"),
])
ms = " · ".join(f"M{k} {nm} ({fmt(t)})" for k, nm, t in MS)
para(doc, "Meilensteine: " + ms + ".", 9.5, after=2, before=4)
# 7
heading(doc, "7.", "Finanzierung")
para(doc, "Für das Projekt entstehen keine nennenswerten Kosten. Benötigt werden digitale Recherchemittel, ein Computer und ein Canva-Zugang für das Layout (vorhanden). Für den Druck von Booklet, Merkblatt und Kartensets: [Anzahl der Ausdrucke und Drucker klären].", 10, after=2)

# SMART
h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(4); h.paragraph_format.keep_with_next = True
run(h, "Unsere zwei SMART-Ziele", 13, True, NAVY)
goal_box(doc, "Ziel 1 – Wissen über den Ausbildungsvertrag",
         "Bis zum 08.10.2026 kennen wir die wichtigsten Regelungen aus Berufsbildungsgesetz und Ausbildungsvertrag und können für jeden der sechs Prüfungspunkte (Beginn und Dauer, tägliche Ausbildungszeit, Vergütung, Urlaub, Probezeit, Kündigung) mindestens ein Beispiel erklären. Das Booklet deckt alle sechs Punkte jeweils mit Aufgaben und Lösung ab.",
         [("Beginn und Dauer", "Verkürzung auf gemeinsamen Antrag von Azubi und Ausbildenden (§ 8 Abs. 1 BBiG)"),
          ("Tägliche Ausbildungszeit", "Jugendliche höchstens 8 Stunden täglich (§ 8 JArbSchG)"),
          ("Vergütung", "mindestens 724 € im 1. Jahr bei Ausbildungsbeginn 2026 (§ 17 BBiG)"),
          ("Urlaub", "ab 18 Jahren mindestens 24 Werktage, darunter 25, 27 oder 30 (§ 3 BUrlG, § 19 JArbSchG)"),
          ("Probezeit", "ein bis vier Monate (§ 20 BBiG)"),
          ("Kündigung", "in der Probezeit jederzeit ohne Frist, aber schriftlich (§ 22 Abs. 1 und 3 BBiG)")])
goal_box(doc, "Ziel 2 – Anwendung im Ausbildungsalltag",
         "Bis zur Präsentation am 08.10.2026 können Mitschüler mit dem Booklet typische Situationen der passenden Regelung zuordnen: Im Abschluss-Quiz mit fünf neuen Situationen erreichen sie im Durchschnitt mindestens 4 von 5 richtige Antworten (80 %).",
         [("Azubi will in der Probezeit kündigen", "jederzeit ohne Frist, schriftlich (§ 22 Abs. 1 und 3 BBiG)"),
          ("16-Jähriger fragt nach Mindesturlaub", "27 Werktage (§ 19 JArbSchG)"),
          ("Betrieb bietet im 1. Jahr (Beginn 2026) 600 € an", "unzulässig, mindestens 724 € (§ 17 BBiG)"),
          ("Azubi will nach der Probezeit den Beruf wechseln", "Kündigung mit vier Wochen Frist, schriftlich (§ 22 Abs. 2 und 3 BBiG)"),
          ("Arbeitstag vor der schriftlichen Abschlussprüfung", "Betrieb muss freistellen (§ 15 BBiG)")])

# Fuss
sec.footer.paragraphs[0].style.font.size = Pt(8); sec.footer.paragraphs[0].style.font.name = FONT
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, "Projektskizze · Yasin & Mido · Seite ", 8, False, "5B6680"); field(fp, "PAGE"); run(fp, " von ", 8, False, "5B6680"); field(fp, "NUMPAGES")

path = os.path.join(OUT, "projektskizze.docx")
doc.save(path)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, path], capture_output=True, timeout=240)
import pymupdf
d = pymupdf.open(os.path.join(OUT, "projektskizze.pdf")); print("Seiten:", len(d))
