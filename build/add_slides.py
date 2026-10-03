# -*- coding: utf-8 -*-
"""Fuegt drei Planungsfolien in praesentation.pptx ein, ohne bestehende Folien zu veraendern."""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

H = lambda s: RGBColor.from_string(s.lstrip("#"))
NAVY, NAVY2, ACC = H("14264B"), H("22386B"), H("F5A800")
ACC_L, BLUE_L, INK = H("FFF1CC"), H("E8EEF9"), H("1A2238")
PPTX = os.path.join(os.path.dirname(__file__), "..", "output", "praesentation.pptx")
IMG = os.path.join(os.path.dirname(__file__), "img")

prs = Presentation(PPTX)
n_before = len(prs.slides)
BLANK = prs.slide_layouts[6]
orig_ids = list(prs.slides._sldIdLst)

def rect(s, x, y, w, h, fill):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background(); r.shadow.inherit = False; return r

def text(s, x, y, w, h, t, size=20, bold=False, color=INK, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05)
    paras = t if isinstance(t, list) else [t]
    for i, pt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i: p.space_before = Pt(size * 0.4)
        runs = pt if isinstance(pt, list) else [(pt, bold, color)]
        for rt, rb, rc in runs:
            r = p.add_run(); r.text = rt; r.font.size = Pt(size); r.font.bold = rb; r.font.color.rgb = rc; r.font.name = "Arial"
    return tb

def base(title, kicker, notes):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 0.35, 7.5, ACC)
    text(s, 0.7, 0.3, 11, 0.35, kicker.upper(), 13, True, H("B07A00"))
    text(s, 0.7, 0.6, 12, 0.9, title, 32, True, NAVY)
    s.notes_slide.notes_text_frame.text = notes
    return s

def card(s, x, y, w, h, head, body, hs=20, bs=20, fill=BLUE_L, bar=NAVY2):
    rect(s, x, y, w, h, fill); rect(s, x, y, 0.1, h, bar)
    text(s, x + 0.25, y + 0.1, w - 0.4, 0.5, head, hs, True, NAVY)
    text(s, x + 0.25, y + 0.1 + hs * 0.026 + 0.14, w - 0.4, h - 0.8, body, bs)

# ---- 1 Projektskizze auf einen Blick
s1 = base("Projektskizze auf einen Blick", "Aufgabe, Problem, Ziel",
          "Yasin: Unsere Projektskizze auf einen Blick: Wir erstellen ein Prüfungsvorbereitungsbooklet für Mitschüler, Laufzeit vom 10. September bis 8. Oktober.\n\n"
          "Mido: Dazu haben wir zwei SMART-Ziele: erstens Wissen über den Ausbildungsvertrag, zweitens Anwendung im Ausbildungsalltag – messbar am Abschluss-Quiz, beide bis zum 8. Oktober.")
card(s1, 0.7, 1.65, 4.0, 1.75, "Ziel", "Booklet zur Prüfungsvorbereitung mit Prüfungstraining", 20, 19)
card(s1, 4.9, 1.65, 4.0, 1.75, "Zielgruppe", "Mitschüler vor der Abschlussprüfung am 25.11.", 20, 19)
card(s1, 9.1, 1.65, 3.55, 1.75, "Laufzeit", "10.09. – 08.10.2026", 20, 22, ACC_L, ACC)
card(s1, 0.7, 3.65, 5.95, 3.4, "SMART-Ziel 1 · Wissen",
     "Bis 08.10. die wichtigsten Regelungen aus BBiG und Ausbildungsvertrag kennen und für alle sechs Prüfungspunkte ein Beispiel erklären können.", 22, 21)
card(s1, 6.85, 3.65, 5.8, 3.4, "SMART-Ziel 2 · Anwendung",
     "Bis 08.10. ordnen Mitschüler mit dem Booklet typische Situationen der passenden Regelung zu: im Abschluss-Quiz im Schnitt mindestens 4 von 5 richtig.", 22, 21, ACC_L, ACC)

# ---- 2 Projektstrukturplan
s2 = base("Projektstrukturplan", "Durchführung",
          "Yasin: Unser Projektstrukturplan zeigt alle Aufgaben des Projekts: fünf Teilaufgaben und zwei einzelne Arbeitspakete, zusammen 19 Arbeitspakete.\n\n"
          "Mido: Jedes Arbeitspaket hat ein klares Ergebnis und eine verantwortliche Person – Yasin, Mido oder wir beide.")
s2.shapes.add_picture(os.path.join(IMG, "psp_folie.png"), Inches(0.75), Inches(1.45), width=Inches(12.1))

# ---- 3 Projektablaufplan
s3 = base("Projektablaufplan (Soll)", "Durchführung",
          "Yasin: Der Projektablaufplan legt für jedes Arbeitspaket Dauer und Zeitraum fest, vom 10. September bis zum 8. Oktober.\n\n"
          "Mido: Fünf Meilensteine sichern den Plan – der letzte ist die Präsentation am 8. Oktober.")
s3.shapes.add_picture(os.path.join(IMG, "gantt_folie.png"), Inches(0.75), Inches(1.45), width=Inches(12.1))

# ---- Reihenfolge: Titel, Projekt in Kuerze, [Skizze, PSP, Gantt], Rest ..., alte Zeitplan-Folie als Reserve ans Ende
lst = prs.slides._sldIdLst
all_ids = list(lst)
new = all_ids[n_before:]
for e in all_ids: lst.remove(e)
order = orig_ids[:2] + new + orig_ids[3:] + [orig_ids[2]]
for e in order: lst.append(e)
prs.save(PPTX)
print("Folien vorher", n_before, "nachher", len(prs.slides))
