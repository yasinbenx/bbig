# -*- coding: utf-8 -*-
"""PSP (PDF/PNG), Projektablaufplan (PDF/XLSX) fuer v4."""
import os, sys, textwrap
sys.path.insert(0, os.path.dirname(__file__))
from html import escape as esc
from playwright.sync_api import sync_playwright
from plan4 import *
OUT = os.path.join(os.path.dirname(__file__), "..", "output", "v4"); os.makedirs(OUT, exist_ok=True)
NAVY, ACC, INK, MUTED = "#14264B", "#F5A800", "#1A2238", "#5B6680"
CSS = f"""@page {{ size: 297mm 210mm; margin: 0 }} *{{box-sizing:border-box;margin:0;padding:0}} html{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
body{{font-family:'Liberation Sans',Arial,sans-serif;color:{INK};font-size:10pt;line-height:1.35}}
.page{{width:297mm;height:210mm;padding:12mm 14mm 10mm 14mm;position:relative;overflow:hidden;page-break-after:always;break-after:page}} .page:last-child{{page-break-after:auto;break-after:auto}}
.k{{color:#B07800;font-size:8.5pt;font-weight:700;letter-spacing:.12em;text-transform:uppercase}} h1{{font-size:19pt;color:{NAVY};line-height:1.15;margin:.5mm 0 1mm}}
.bar{{width:24mm;height:1.5mm;background:{ACC};margin:1.5mm 0 4mm}} .sub{{color:{MUTED};font-size:9.5pt;margin-bottom:3mm}}
.foot{{position:absolute;left:14mm;right:14mm;bottom:5mm;font-size:8pt;color:{MUTED};display:flex;justify-content:space-between;border-top:.3mm solid #D5D9E3;padding-top:1.2mm}}
table{{border-collapse:collapse;width:100%}} th{{background:{NAVY};color:#fff;text-align:left;padding:1.6mm 2.4mm;font-size:9pt}} td{{padding:0.9mm 2.4mm;border-bottom:.25mm solid #D5D9E3;font-size:9.4pt;vertical-align:top}}
tr.g td{{background:#E8EEF9;font-weight:700;color:{NAVY}}} td.n{{font-weight:700;color:{NAVY};white-space:nowrap}} tr.b td{{background:#FFF1CC}}"""
FOOT = f'<div class="foot"><span>{esc(TITEL)} · Yasin &amp; Mido · Fach GP</span><span>%s</span></div>'
def head(k, t, sub=""): return f'<div class="k">{k}</div><h1>{t}</h1><div class="bar"></div>' + (f'<div class="sub">{sub}</div>' if sub else "")
COL = {"Y": NAVY, "M": ACC, "YM": "url(#hatch)"}

# ------------------------------------------------------------------ PSP
def psp_svg():
    W, H = 1123, 600; cw, gap, x0 = 200, 14, 33
    cx = lambda i: x0 + i * (cw + gap)
    o = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" width="100%" font-family="Liberation Sans, Arial"><defs><pattern id="hatch" width="8" height="8" patternUnits="userSpaceOnUse" patternTransform="rotate(45)"><rect width="8" height="8" fill="{NAVY}"/><rect width="4" height="8" fill="{ACC}"/></pattern></defs>']
    o.append(f'<rect x="{W/2-210}" y="6" width="420" height="50" rx="6" fill="{NAVY}"/><text x="{W/2}" y="26" text-anchor="middle" fill="{ACC}" font-size="11" font-weight="700" letter-spacing="1.5">PROJEKT</text><text x="{W/2}" y="44" text-anchor="middle" fill="#fff" font-size="15" font-weight="700">{esc(TITEL)}</text>')
    o.append(f'<path d="M{W/2} 56 V78 M{cx(0)+cw/2} 78 H{cx(4)+cw/2}" stroke="{NAVY}" stroke-width="2" fill="none"/>')
    ymax = 0
    for i in range(1, 6):
        x = cx(i - 1); xm = x + cw / 2
        o.append(f'<path d="M{xm} 78 V96" stroke="{NAVY}" stroke-width="2"/>')
        t = textwrap.wrap(TA[i], 24)
        o.append(f'<rect x="{x}" y="96" width="{cw}" height="64" rx="6" fill="{ACC}"/><text x="{x+12}" y="118" font-size="20" font-weight="700" fill="{NAVY}">{i}</text>')
        for k, line in enumerate(t[:3]): o.append(f'<text x="{x+36}" y="{116+k*15}" font-size="12.5" font-weight="700" fill="{NAVY}">{esc(line)}</text>')
        y = 172; lx = x + 14
        aps = [a for a in AP if teil(a[0]) == i]
        last = y; lines_o = len(o); o.append("")
        for n, nm, d, pr, who, _ in aps:
            lines = textwrap.wrap(nm, 27); h = 20 + len(lines) * 13.5 + 16
            o.append(f'<path d="M{x+10} {y+h/2} H{x+24}" stroke="{NAVY}" stroke-width="1.5"/>')
            o.append(f'<rect x="{x+24}" y="{y}" width="{cw-24}" height="{h}" rx="5" fill="#fff" stroke="{NAVY}" stroke-width="1.6"/><rect x="{x+24}" y="{y}" width="5" height="{h}" rx="2" fill="{COL[who]}"/>')
            o.append(f'<text x="{x+36}" y="{y+16}" font-size="11.5" font-weight="700" fill="{NAVY}">{n}</text>')
            for k, line in enumerate(lines): o.append(f'<text x="{x+36}" y="{y+31+k*13.5}" font-size="11" fill="{INK}">{esc(line)}</text>')
            dd = f"{DUR[n]} AT" if DUR[n] else "nach 08.10."
            o.append(f'<text x="{x+36}" y="{y+h-6}" font-size="9.5" fill="{MUTED}">{OWNER[who]} · {dd}</text>')
            last = y + h; y += h + 10
        o[lines_o] = f'<path d="M{x+10} 160 V{last-(h/2)}" stroke="{NAVY}" stroke-width="2" fill="none"/>'
        ymax = max(ymax, y)
        if i < 5: o.append(f'<path d="M{x+cw+2} 128 h{gap-4}" stroke="{ACC}" stroke-width="0"/>')
    wy = ymax + 24
    o.append(f'<text x="{x0}" y="{wy}" font-size="11" font-weight="700" fill="{MUTED}" letter-spacing="1">WASSERFALLMODELL</text>')
    for i in range(5):
        xx = cx(i)
        o.append(f'<path d="M{xx} {wy+8} h{cw-14} l14 17 l-14 17 h-{cw-14} l14 -17z" fill="{NAVY if i%2==0 else "#22386B"}"/><text x="{xx+cw/2}" y="{wy+30}" text-anchor="middle" fill="#fff" font-size="12" font-weight="700">{i+1}  {esc(["Planung","Recherche","Booklet","Präsentation","Abschluss"][i])}</text>')
    o.append("</svg>"); return "\n".join(o), ymax
svg, ym = psp_svg()
def psp_pages():
    rows = ""
    for i in range(1, 6):
        rows += f'<tr class="g"><td class="n">{i}</td><td colspan="4">{esc(TA[i])}</td></tr>'
        for n, nm, d, pr, who, _ in [a for a in AP if teil(a[0]) == i]:
            rows += f'<tr><td class="n">{n}</td><td>{esc(nm)}</td><td>{OWNER[who]}</td><td>{dur_txt(n)}</td><td>{esc(pred_txt(n))}</td></tr>'
    p1 = f'<section class="page">{head("Projektstrukturplan", esc(TITEL), "Wasserfallmodell: Die fünf Teilaufgaben bauen von links nach rechts aufeinander auf; Arbeitspakete (AP) sind nach Teilaufgabe nummeriert (1.1, 1.2 …). Randbalken der Arbeitspakete: Navy = Yasin · Orange = Mido · gestreift = beide.")}{svg}{FOOT % "Seite 1 von 2"}</section>'
    p2 = f'<section class="page">{head("Projektstrukturplan · Arbeitspaketliste", "Teilaufgaben und Arbeitspakete im Überblick")}<table><tr><th style="width:14mm">Nr.</th><th>Teilaufgabe / Arbeitspaket</th><th style="width:38mm">Verantwortlich</th><th style="width:22mm">Dauer (AT)</th><th style="width:30mm">Vorgänger</th></tr>{rows}</table><p style="font-size:8.8pt;color:{MUTED};margin-top:3mm">AT = Arbeitstage (Montag bis Freitag). AP 5.3 findet nach der Präsentation statt; der Termin richtet sich nach der Vorgabe der Lehrkraft. Termine siehe Projektablaufplan.</p>{FOOT % "Seite 2 von 2"}</section>'
    return p1 + p2

# ------------------------------------------------------------------ Gantt + Tabelle
def gantt_svg():
    W = 1123 - 0; lab = 300; cw = 34; x0 = 20; gx = x0 + lab; rh = 27; gh = 22
    rows = []; y = 66
    for i in range(1, 6):
        rows.append(("g", i, y)); y += gh
        for a in [a for a in AP if teil(a[0]) == i]: rows.append(("a", a, y)); y += rh
    H = y + 10
    o = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" width="100%" font-family="Liberation Sans, Arial"><defs><pattern id="hatch" width="8" height="8" patternUnits="userSpaceOnUse" patternTransform="rotate(45)"><rect width="8" height="8" fill="{NAVY}"/><rect width="4" height="8" fill="{ACC}"/></pattern></defs>']
    # Kopf: Wochen (KW) und Tage
    cols = []
    for i, d in enumerate(WORKDAYS): cols.append(gx + i * cw)
    wk = {}
    for i, d in enumerate(WORKDAYS): wk.setdefault(d.isocalendar()[1], []).append(i)
    for w, idx in wk.items():
        xa = gx + idx[0] * cw; xb = gx + (idx[-1] + 1) * cw
        o.append(f'<rect x="{xa}" y="4" width="{xb-xa}" height="22" fill="{NAVY}" stroke="#fff" stroke-width="1"/><text x="{(xa+xb)/2}" y="19" text-anchor="middle" fill="#fff" font-size="11" font-weight="700">KW {w}</text>')
    tag = ["Mo", "Di", "Mi", "Do", "Fr"]
    for i, d in enumerate(WORKDAYS):
        xa = gx + i * cw
        o.append(f'<rect x="{xa}" y="26" width="{cw}" height="34" fill="#F6F8FC" stroke="#D5D9E3"/><text x="{xa+cw/2}" y="40" text-anchor="middle" font-size="9" fill="{MUTED}">{tag[d.weekday()]}</text><text x="{xa+cw/2}" y="54" text-anchor="middle" font-size="10.5" font-weight="700" fill="{NAVY}">{d.day:02d}.{d.month:02d}.</text>' if False else
                 f'<rect x="{xa}" y="26" width="{cw}" height="34" fill="#F6F8FC" stroke="#D5D9E3"/><text x="{xa+cw/2}" y="39" text-anchor="middle" font-size="9" fill="{MUTED}">{tag[d.weekday()]}</text><text x="{xa+cw/2}" y="53" text-anchor="middle" font-size="9.5" font-weight="700" fill="{NAVY}">{d.day}.{d.month}.</text>')
    o.append(f'<text x="{x0+4}" y="52" font-size="10" font-weight="700" fill="{NAVY}">Arbeitspaket</text>')
    # Zeilen
    for kind, v, yy in rows:
        if kind == "g":
            o.append(f'<rect x="{x0}" y="{yy}" width="{lab+21*cw}" height="{gh}" fill="#E8EEF9"/><text x="{x0+6}" y="{yy+13.5}" font-size="10.5" font-weight="700" fill="{NAVY}">{v}  {esc(TA[v])}</text>')
        else:
            n, nm, d, pr, who, _ = v
            short = nm if len(nm) <= 50 else nm[:47].rsplit(" ", 1)[0] + " …"
            o.append(f'<text x="{x0+6}" y="{yy+15}" font-size="10" fill="{INK}"><tspan font-weight="700" fill="{NAVY}">{n}</tspan>  {esc(short)}</text>')
            o.append(f'<line x1="{x0}" x2="{gx+21*cw}" y1="{yy+rh}" y2="{yy+rh}" stroke="#E3E7F0"/>')
            if n in ES:
                xa = gx + ES[n] * cw + 2; wd = (EF[n] - ES[n] + 1) * cw - 4
                o.append(f'<rect x="{xa}" y="{yy+4}" width="{wd}" height="{rh-8}" rx="3" fill="{COL[who]}" stroke="{NAVY}" stroke-width="1"/>')
            else:
                o.append(f'<text x="{gx+21*cw-4}" y="{yy+15}" text-anchor="end" font-size="9.5" fill="{MUTED}">→ nach der Präsentation (Termin laut Lehrkraft)</text>')
    # vertikale Linien (Woche)
    for w, idx in wk.items():
        o.append(f'<line x1="{gx+idx[0]*cw}" x2="{gx+idx[0]*cw}" y1="26" y2="{y}" stroke="#9AA3B8" stroke-width="1.2"/>')
    o.append(f'<line x1="{gx+21*cw}" x2="{gx+21*cw}" y1="26" y2="{y}" stroke="#9AA3B8" stroke-width="1.2"/>')
    # Meilensteine
    for k, nm, d in MS:
        i = WORKDAYS.index(d); xm = gx + (i + 1) * cw - 3
        o.append(f'<line x1="{xm}" x2="{xm}" y1="60" y2="{y}" stroke="#C0392B" stroke-width="1" stroke-dasharray="3,3" opacity=".7"/><path d="M{xm} 61 l6 6 l-6 6 l-6 -6z" fill="#C0392B"/>')
        o.append(f'<text x="{xm}" y="{y+9}" text-anchor="middle" font-size="9.5" font-weight="700" fill="#C0392B">{k}</text>')
    o.append("</svg>")
    return "\n".join(o), H
gs, gH = gantt_svg()
def ablauf_pages():
    rows = ""
    for i in range(1, 6):
        rows += f'<tr class="g"><td class="n">{i}</td><td colspan="6">{esc(TA[i])}</td></tr>'
        for n, nm, d, pr, who, _ in [a for a in AP if teil(a[0]) == i]:
            rows += f'<tr><td class="n">{n}</td><td>{esc(nm)}</td><td>{OWNER[who]}</td><td style="text-align:center">{dur_txt(n)}</td><td>{start_txt(n)}</td><td>{end_txt(n)}</td><td>{esc(pred_txt(n))}</td></tr>'
    p1 = f'<section class="page">{head("Projektablaufplan", "Ablauftabelle mit Soll-Terminen", "Laufzeit 10.09. bis 08.10.2026 · 21 Arbeitstage (Montag bis Freitag) · Wasserfallmodell")}<table><tr><th style="width:13mm">AP</th><th>Arbeitspaket</th><th style="width:34mm">Verantwortlich</th><th style="width:20mm;text-align:center">Dauer (AT)</th><th style="width:31mm">Start</th><th style="width:38mm">Ende</th><th style="width:24mm">Vorgänger</th></tr>{rows}</table>{FOOT % "Seite 1 von 3"}</section>'
    leg = "".join(f'<tr><td class="n" style="color:#C0392B">{k}</td><td>{esc(nm)}</td><td>{fmt(d, True)}</td></tr>' for k, nm, d in MS)
    p2 = f'<section class="page">{head("Projektablaufplan", "Gantt-Diagramm", "Balken: Navy = Yasin · Orange = Mido · gestreift = beide · rote Raute = Meilenstein (Ende des Tages)")}{gs}{FOOT % "Seite 2 von 3"}</section>'
    crit = ", ".join(CRIT)
    p3 = f'<section class="page">{head("Projektablaufplan", "Meilensteine und kritischer Weg")}<table style="width:190mm"><tr><th style="width:14mm">Nr.</th><th>Meilenstein</th><th style="width:34mm">Termin</th></tr>{leg}</table><p style="margin-top:6mm;font-size:10pt"><b>Kritischer Weg</b> (kein Puffer, Verzug verschiebt die Präsentation): {crit}.</p><p style="margin-top:2mm;font-size:10pt"><b>Puffer:</b> AP 3.3 hat {SLACK["3.3"]} Arbeitstage, AP 4.2 hat {SLACK["4.2"]} Arbeitstag Puffer.</p><p style="margin-top:2mm;font-size:9.5pt;color:{MUTED}">Die Termine sind Soll-Termine aus Dauer, Vorgängern und frühestem Start (AP 4.2 frühestens am 02.10.2026). Tatsächliche Termine (Ist) werden im Projektbericht gegenübergestellt.</p>{FOOT % "Seite 3 von 3"}</section>'
    return p1 + p2 + p3
def render(body, name, png=None):
    html = f'<!doctype html><html lang="de"><head><meta charset="utf-8"><title>{name}</title><style>{CSS}</style></head><body>{body}</body></html>'
    hp = os.path.join(OUT, "_" + name + ".html"); open(hp, "w", encoding="utf-8").write(html)
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome", args=["--no-sandbox"])
        pg = b.new_page(viewport={"width": 1123, "height": 794}, device_scale_factor=2); pg.goto("file://" + os.path.abspath(hp)); pg.wait_for_timeout(300)
        over = pg.evaluate("() => [...document.querySelectorAll('.page')].map((s,i)=>{let m=0;s.querySelectorAll('*').forEach(e=>{const r=e.getBoundingClientRect(); if(r.height>0&&!e.closest('.foot')) m=Math.max(m,r.bottom-s.getBoundingClientRect().top)}); return [i+1, Math.round(m), s.clientHeight]})")
        for i, m, h in over:
            if m > h - 38: print("WARN", name, "Seite", i, "Inhalt bis", m, "von", h)
        pg.pdf(path=os.path.join(OUT, name + ".pdf"), prefer_css_page_size=True, print_background=True)
        if png: pg.locator(".page").first.screenshot(path=os.path.join(OUT, png))
        b.close()
    os.remove(hp)
render(psp_pages(), "projektstrukturplan_v4", "projektstrukturplan_v4.png")
render(ablauf_pages(), "projektablaufplan_v4")

# ------------------------------------------------------------------ XLSX
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
wb = Workbook(); ws = wb.active; ws.title = "Ablaufplan"
hf = PatternFill("solid", fgColor="14264B"); gf = PatternFill("solid", fgColor="E8EEF9"); thin = Side(style="thin", color="BFC6D6"); bd = Border(bottom=thin)
ws["A1"] = f"Projektablaufplan – {TITEL}"; ws["A1"].font = Font(bold=True, size=14, color="14264B")
ws["A2"] = "Laufzeit 10.09.2026 bis 08.10.2026 · 21 Arbeitstage (Mo–Fr) · Wasserfallmodell · Soll-Termine"; ws["A2"].font = Font(size=10, color="5B6680")
hdr = ["AP", "Arbeitspaket", "Verantwortlich", "Dauer (AT)", "Start", "Ende", "Vorgänger"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(4, j, h); c.fill = hf; c.font = Font(bold=True, color="FFFFFF")
r = 5
for i in range(1, 6):
    for j in range(1, 8): ws.cell(r, j).fill = gf
    ws.cell(r, 1, i).font = Font(bold=True, color="14264B"); ws.cell(r, 2, TA[i]).font = Font(bold=True, color="14264B"); r += 1
    for n, nm, d, pr, who, _ in [a for a in AP if teil(a[0]) == i]:
        vals = [n, nm, OWNER[who], d if d else "–", start(n) if n in ES else "nach 08.10.2026", end(n) if n in EF else "[laut Lehrkraft]", pred_txt(n)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.border = bd
            if isinstance(v, dt.date): c.number_format = "DD.MM.YYYY"
        r += 1
for col, w in zip("ABCDEFG", [8, 62, 18, 11, 18, 24, 14]): ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
# Gantt
g = wb.create_sheet("Gantt"); g["A1"] = "Gantt-Diagramm (Arbeitstage)"; g["A1"].font = Font(bold=True, size=14, color="14264B")
g.cell(3, 1, "AP").fill = hf; g.cell(3, 2, "Arbeitspaket").fill = hf; g.cell(3, 3, "Verantw.").fill = hf
for c in (1, 2, 3): g.cell(3, c).font = Font(bold=True, color="FFFFFF")
tag = ["Mo", "Di", "Mi", "Do", "Fr"]
for i, d in enumerate(WORKDAYS):
    c = g.cell(3, 4 + i, f"{tag[d.weekday()]} {d.day}.{d.month}."); c.fill = hf; c.font = Font(bold=True, color="FFFFFF", size=8); c.alignment = Alignment(horizontal="center", text_rotation=90)
    g.column_dimensions[get_column_letter(4 + i)].width = 4.2
g.row_dimensions[3].height = 48
fill = {"Y": "14264B", "M": "F5A800", "YM": "6B7FB5"}
r = 4
for i in range(1, 6):
    for j in range(1, 4 + 21): g.cell(r, j).fill = gf
    g.cell(r, 1, i).font = Font(bold=True, color="14264B"); g.cell(r, 2, TA[i]).font = Font(bold=True, color="14264B"); r += 1
    for n, nm, d, pr, who, _ in [a for a in AP if teil(a[0]) == i]:
        g.cell(r, 1, n); g.cell(r, 2, nm); g.cell(r, 3, OWNER[who])
        if n in ES:
            for k in range(ES[n], EF[n] + 1): g.cell(r, 4 + k).fill = PatternFill("solid", fgColor=fill[who])
        else: g.cell(r, 4, "nach der Präsentation (Termin laut Lehrkraft)").font = Font(italic=True, color="5B6680")
        r += 1
r += 1; g.cell(r, 2, "Meilensteine").font = Font(bold=True, color="14264B"); r += 1
for k, nm, d in MS:
    g.cell(r, 1, k).font = Font(bold=True, color="C0392B"); g.cell(r, 2, nm); c = g.cell(r, 3, d); c.number_format = "DD.MM.YYYY"
    g.cell(r, 4 + WORKDAYS.index(d)).fill = PatternFill("solid", fgColor="C0392B"); r += 1
r += 1; g.cell(r, 2, "Legende: dunkelblau = Yasin · orange = Mido · blau = Yasin & Mido · rot = Meilenstein").font = Font(size=9, color="5B6680")
g.column_dimensions["A"].width = 6; g.column_dimensions["B"].width = 62; g.column_dimensions["C"].width = 16; g.freeze_panes = "D4"
m = wb.create_sheet("Meilensteine"); m.append(["Nr.", "Meilenstein", "Termin"])
for c in m[1]: c.fill = hf; c.font = Font(bold=True, color="FFFFFF")
for k, nm, d in MS: m.append([k, nm, d]); m.cell(m.max_row, 3).number_format = "DD.MM.YYYY"
m.column_dimensions["B"].width = 50; m.column_dimensions["C"].width = 14
for ws_ in (ws, g, m):
    ws_.page_setup.orientation = "landscape"; ws_.page_setup.fitToWidth = 1; ws_.sheet_properties.pageSetUpPr.fitToPage = True; ws_.page_setup.fitToHeight = 0
wb.save(os.path.join(OUT, "projektablaufplan_v4.xlsx"))
print("ok")
