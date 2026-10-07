# -*- coding: utf-8 -*-
"""Bausteine fuer die unterrichtsreife Praesentation (v3): Layouts, echte PowerPoint-Animationen, Notizen."""
import os, re, json
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
IC = lambda n: os.path.join(os.path.dirname(__file__), "img", "icons", n + ".png")
QR = os.path.join(ROOT, "output/v2/qr/qr.png")
URL_KURZ = "yasinbenx.github.io/azubi-vertrag"
FONT = "Calibri"

H = lambda s: RGBColor.from_string(s.lstrip("#"))
NAVY, NAVY2, ACC = H("1F2A44"), H("34456B"), H("F5A623")
ACC_L, LIGHT, GREY, INK, MUTED, WHITE, LINE = H("FFF0D1"), H("EEF1F8"), H("E4E8F1"), H("1F2A44"), H("4A5568"), H("FFFFFF"), H("B8C0D0")
AMBER, GREEN, GREEN_L, RED, RED_L, ORANGE = H("7A4B00"), H("1E7B4F"), H("E3F3EA"), H("B3261E"), H("FCE6E3"), H("B85C00")
ORANGE_L = H("FCEBDD")

PHASES = ["Einstieg", "Block 1", "Block 2", "Block 3", "Üben", "Besprechen", "Abschluss"]
PH_TARGET = {"Einstieg": 300, "Block 1": 330, "Block 2": 330, "Block 3": 300, "Üben": 720, "Besprechen": 420, "Abschluss": 300}

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
META = []


def mmss(sec): return f"{sec // 60}:{sec % 60:02d}"


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5, name=None):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb = line; r.line.width = Pt(lw)
    r.shadow.inherit = False
    if name: r.name = name
    return r


def text(s, x, y, w, h, t, size=28, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None, name=None, space=0.3, font=FONT):
    tb = shape or s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, pt in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        if i: p.space_before = Pt(size * space)
        for rt, rb, rc in (pt if isinstance(pt, list) else [(pt, bold, color)]):
            r = p.add_run(); r.text = rt; r.font.size = Pt(size); r.font.bold = rb; r.font.color.rgb = rc; r.font.name = font
    if name: tb.name = name
    return tb


def pic(s, path, x, y, w, h=None, alt="", name=None):
    p = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h) if h else None)
    p._element.nvPicPr.cNvPr.set("descr", alt)
    if name: p.name = name
    return p


def circle(s, x, y, d, label, size=28, fill=NAVY, fg=WHITE, name=None):
    c = rect(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL, name=name)
    text(s, 0, 0, 0, 0, str(label), size, True, fg, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, shape=c)
    c.text_frame.margin_left = c.text_frame.margin_right = 0; c.text_frame.word_wrap = False
    return c


def card(s, x, y, w, h, head, body, hs=28, bs=28, fill=LIGHT, bar=NAVY2, name=None):
    r = rect(s, x, y, w, h, fill, name=name); rect(s, x, y, 0.12, h, bar)
    yy = y + 0.08
    if head:
        text(s, x + 0.3, yy, w - 0.45, 0.65, head, hs, True, NAVY); yy += hs / 72 * 1.35
    text(s, x + 0.3, yy, w - 0.45, h - (yy - y) - 0.05, body, bs)
    return r


def qr_block(s, x, y, size, alt="QR-Code zur Übungsseite"):
    pad = size * 0.05
    rect(s, x, y, size + 2 * pad, size + 2 * pad, WHITE, LINE)
    pic(s, QR, x + pad, y + pad, size, size, alt)


def band(s, idx):
    w = 13.333 / len(PHASES)
    for i, ph in enumerate(PHASES):
        cur, done = i == idx, i < idx
        r = rect(s, i * w, 0, w, 0.5, ACC if cur else (NAVY2 if done else NAVY), WHITE, lw=1)
        text(s, i * w, 0, w, 0.5, ph, 24, cur, NAVY if cur else WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def badge(s, t, x=0.55, w=None, fill=NAVY, fg=WHITE, y=6.92):
    w = w or (0.45 + len(t) * 0.17)
    rect(s, x, y, w, 0.46, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.4
    text(s, x, y, w, 0.46, t, 24, True, fg, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return x + w + 0.15


def para_corner(s, t):
    if t: text(s, 8.3, 6.92, 4.5, 0.46, t, 24, False, MUTED, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- Animation (echtes p:timing)
NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
P = lambda tag: f"{{{NS}}}{tag}"


class Timing:
    def __init__(self):
        self.id = 2; self.blocks = []; self.bld = []

    def nid(self): self.id += 1; return self.id

    def _effect(self, spid, node):
        a, b = self.nid(), self.nid()
        return (f'<p:par><p:cTn id="{a}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')

    def click(self, shapes):
        outer, inner = self.nid(), self.nid()
        eff = "".join(self._effect(sh.shape_id, "clickEffect" if i == 0 else "withEffect") for i, sh in enumerate(shapes))
        self.blocks.append(f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                           f'<p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{eff}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
        self._bld(shapes)

    def auto(self, shapes, step_ms):
        """startet automatisch beim Folienwechsel; jede Form erscheint nach step_ms (Countdown)"""
        outer = self.nid(); inner = []
        for i, sh in enumerate(shapes):
            pid = self.nid()
            inner.append(f'<p:par><p:cTn id="{pid}" fill="hold"><p:stCondLst><p:cond delay="{i * step_ms}"/></p:stCondLst><p:childTnLst>{self._effect(sh.shape_id, "afterEffect")}</p:childTnLst></p:cTn></p:par>')
        self.blocks.append(f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{"".join(inner)}</p:childTnLst></p:cTn></p:par>')
        self._bld(shapes)

    def _bld(self, shapes):
        for sh in shapes:
            if sh.shape_type is not None and sh._element.tag == P("sp") and sh.shape_id not in [x for x in self.bld]:
                self.bld.append(sh.shape_id)

    def attach(self, slide):
        bld = "".join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for i in self.bld)
        xml = (f'<p:timing xmlns:p="{NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
               f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(self.blocks)}</p:childTnLst></p:cTn>'
               f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
               f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
               f'<p:bldLst>{bld}</p:bldLst></p:timing>')
        slide._element.append(etree.fromstring(xml))


def animate(s, groups=None, auto=None, step_ms=1000):
    """groups: Liste von Klick-Schritten (je Liste von Shapes); auto: Shapes, die automatisch nacheinander erscheinen"""
    t = Timing()
    if auto: t.auto(auto, step_ms)
    for g in (groups or []): t.click(g if isinstance(g, list) else [g])
    t.attach(s)
    s._anim = dict(auto=[a.name for a in (auto or [])], groups=[[x.name for x in (g if isinstance(g, list) else [g])] for g in (groups or [])])


# ---------------------------------------------------------------- Folienfabrik
def new_slide(phase, kind, title, dauer, spr, sozial, impuls="", erw="", fehl="", med="", puffer="", hidden=False, forbid=None,
              dark=False, anim_note="", ph_idx=None, label=None):
    s = prs.slides.add_slide(BLANK)
    if dark:
        rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0, 0, 0.35, 7.5, ACC)
    else:
        rect(s, 0, 0, 13.333, 7.5, WHITE)
        band(s, PHASES.index(phase) if ph_idx is None else ph_idx)
    s.shapes[0].name = "Hintergrund"
    n = [f"PHASE: {phase} · FOLIENTYP: {kind}",
         f"SPRECHER: {spr}", f"DAUER: {mmss(dauer) if dauer is not None else 'Reserve (nicht in den 45 Minuten)'}", f"SOZIALFORM: {sozial}"]
    if impuls: n.append("IMPULS / MODERATION: " + impuls)
    if erw: n.append("ERWARTETE ANTWORTEN: " + erw)
    if fehl: n.append("FEHLVORSTELLUNGEN UND REAKTION: " + fehl)
    if med: n.append("MEDIEN / TAFEL: " + med)
    if anim_note: n.append("ANIMATION (Klickfolge): " + anim_note)
    if puffer: n.append("ZEITPUFFER / STREICHEN: " + puffer)
    s.notes_slide.notes_text_frame.text = "\n".join(n)
    if hidden: s._element.set("show", "0")
    s._meta = dict(nr=len(prs.slides), phase=phase, kind=kind, title=title or label or kind, dauer=dauer, spr=spr, sozial=sozial, impuls=impuls, erw=erw,
                   fehl=fehl, med=med, puffer=puffer, hidden=hidden, forbid=forbid or [], pair=None)
    META.append(s._meta)
    return s


def titel(s, t, size=40):
    text(s, 0.6, 0.58, 12.2, 0.95, t, size, True, NAVY, anchor=MSO_ANCHOR.MIDDLE, name="Titel")
    rect(s, 0.7, 1.5, 1.6, 0.07, ACC)


def frage_slide(phase, label, frage, dauer, spr, sozial, options=None, forbid=None, size=40, opt_size=28, hidden=False, **kw):
    """Frage-Folie: Frage zentriert, Fragezeichen-Icon, Sozialform-Badge, KEINE Antwort"""
    s = new_slide(phase, "Frage", None, dauer, spr, sozial, forbid=forbid, hidden=hidden, label=label, **kw)
    pic(s, IC("question"), 0.7, 1.0 if not options else 0.8, 1.35, alt="Fragezeichen-Symbol")
    h = 3.9 if not options else 2.0
    text(s, 2.3, 0.75 if options else 1.0, 10.5, h if not options else 1.9, frage, size, True, NAVY, anchor=MSO_ANCHOR.MIDDLE, name="Frage")
    if options:
        cols = 2; cw, ch = 6.0, 1.3
        for i, o in enumerate(options):
            x, y = 0.7 + (i % cols) * (cw + 0.2), 3.0 + (i // cols) * (ch + 0.2)
            rect(s, x, y, cw, ch, WHITE, NAVY, MSO_SHAPE.ROUNDED_RECTANGLE, 2.5).adjustments[0] = 0.12
            circle(s, x + 0.2, y + (ch - 0.75) / 2, 0.75, "ABCD"[i], 28)
            text(s, x + 1.15, y, cw - 1.3, ch, o, opt_size, False, INK, anchor=MSO_ANCHOR.MIDDLE, name=f"Option_{'ABCD'[i]}")
    badge(s, sozial.split(" · ")[0] if len(sozial) < 40 else "Plenum")
    return s


def fall_frage_slide(phase, label, f, dauer, spr, forbid=None, hidden=False, **kw):
    s = new_slide(phase, "Frage", None, dauer, spr, "Plenum · Abstimmung A–D per Handzeichen", forbid=forbid, hidden=hidden, label=label, **kw)
    text(s, 0.6, 0.6, 12.2, 0.6, f["label"], 28, True, AMBER, name="Fallbezeichnung")
    rect(s, 0.7, 1.25, 11.9, 1.75, LIGHT); rect(s, 0.7, 1.25, 0.12, 1.75, ACC)
    text(s, 0.95, 1.25, 11.55, 1.75, f["q"], 26, True, NAVY, anchor=MSO_ANCHOR.MIDDLE, name="Fall")
    for i, a in enumerate(f["a"]):
        y = 3.15 + i * 0.92
        rect(s, 0.7, y, 11.9, 0.82, WHITE, NAVY, MSO_SHAPE.ROUNDED_RECTANGLE, 2).adjustments[0] = 0.15
        circle(s, 0.85, y + 0.09, 0.64, "ABCD"[i], 24)
        text(s, 1.65, y, 10.85, 0.82, a, 24, False, INK, anchor=MSO_ANCHOR.MIDDLE, name=f"Option_{'ABCD'[i]}")
    badge(s, "Abstimmung A–D")
    return s


def denk_slide(phase, secs, spr, sozial="Think-Pair-Share · Partnergespräch", impuls="", **kw):
    s = new_slide(phase, "Denkpause", None, secs, spr, sozial, impuls=impuls or "Stille aushalten; danach Partner:in sprechen lassen. Bei Schweigen: Zeit verlängern, nicht auflösen.", dark=True,
                  anim_note="Countdown-Balken läuft automatisch beim Folienwechsel (Segmente erscheinen nacheinander).", label="Denkpause", **kw)
    rect(s, 5.75, 0.8, 1.8, 1.8, WHITE, shape=MSO_SHAPE.OVAL, name="Icon_Hintergrund")
    pic(s, IC("people"), 6.05, 1.1, 1.2, alt="Zwei Personen im Gespräch")
    text(s, 1, 2.5, 11.3, 1.2, "Denkpause", 66, True, WHITE, PP_ALIGN.CENTER, name="Denkpause")
    text(s, 1, 3.7, 11.3, 1.0, f"{secs} Sekunden – sprecht kurz mit eurer Sitznachbarin / eurem Sitznachbarn.", 30, False, H("DCE3F2"), PP_ALIGN.CENTER)
    n = 10; segs = []; w = 11.0 / n
    for i in range(n):
        segs.append(rect(s, 1.15 + i * w, 5.2, w - 0.1, 0.55, ACC, name=f"Countdown_{i+1}"))
    text(s, 1, 5.9, 11.3, 0.8, f"Countdown: {secs} Sekunden", 28, True, ACC, PP_ALIGN.CENTER)
    animate(s, auto=segs, step_ms=int(secs * 1000 / n))
    return s


def antwort_slide(phase, label, title, dauer, spr, sozial="Plenum", pair=None, **kw):
    s = new_slide(phase, "Antwort", title, dauer, spr, sozial, label=label, **kw)
    rect(s, 0.45, 0.62, 12.43, 6.2, None, ACC, lw=6, name="Antwortrahmen")
    pic(s, IC("check"), 12.0, 0.7, 0.75, alt="Haken: Antwortfolie")
    text(s, 0.7, 0.65, 11.2, 0.9, title, 40, True, NAVY, anchor=MSO_ANCHOR.MIDDLE, name="Titel")
    if pair is not None: pair._meta["pair"] = s._meta["nr"]
    return s


def erklaer_slide(phase, label, title, dauer, spr, sozial="Plenum · Zuhören, Mitschreiben", para=None, **kw):
    s = new_slide(phase, "Erklärfolie", title, dauer, spr, sozial, label=label, **kw)
    titel(s, title); para_corner(s, para)
    return s
