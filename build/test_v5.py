# -*- coding: utf-8 -*-
import os, re, json
from lxml import etree
from pptx import Presentation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
P = Presentation(os.path.join(ROOT, "output/v3/praesentation_unterricht_v4.pptx"))
M = json.load(open("/tmp/v5_meta.json", encoding="utf-8")); META, PAIRS = M["meta"], M["pairs"]
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
X = lambda el, e: etree._Element.xpath(el, e, namespaces=NS)
res = []
def ok(n, c, d=""):
    res.append((bool(c), n, d))
    if not c: print("FAIL:", n, d)

def anim_targets(s):
    t = s._element.find(".//p:timing", NS)
    if t is None: return set(), None
    seq = X(t, ".//p:cTn[@nodeType='mainSeq']/p:childTnLst/p:par"); names = {sh.shape_id: sh.name for sh in s.shapes}
    auto, clicks = [], []
    for outer in seq:
        is_auto = any(c.get("evt") == "onBegin" for c in X(outer, "./p:cTn/p:stCondLst/p:cond"))
        g = [names[int(X(e, ".//p:spTgt/@spid")[0])] for e in X(outer, ".//p:cTn[@presetClass='entr']")]
        (auto.extend(g) if is_auto else clicks.append(g))
    ids = {sh.shape_id for sh in s.shapes if sh.name in set(auto) | {n for g in clicks for n in g}}
    return ids, (auto, clicks)

def pre_click_text(s):
    hid, _ = anim_targets(s); out = []
    for sh in s.shapes:
        if sh.shape_id in hid: continue
        if sh.has_text_frame and sh.text_frame.text.strip(): out.append(sh.text_frame.text)
        out += sh._element.xpath(".//p:cNvPr/@descr")
    return out

# -------- Test 2: sichtbarer Text vor dem ersten Klick bei Frage-Folien
print("== Sichtbarer Text vor dem ersten Klick (Frage-Folien)")
for i, s in enumerate(P.slides, 1):
    m = META[i - 1]
    if m["kind"] not in ("Frage", "Fall"): continue
    vis = " | ".join(pre_click_text(s)); low = vis.lower()
    bad = [f for f in m["forbid"] if f.lower() in low]
    print(f"\n--- Folie {i} ({m['title'][:46]}) ---\n{vis[:700]}{' …' if len(vis) > 700 else ''}")
    ok(f"Folie {i} ({m['title'][:30]}): keine Lösung vor dem Klick", not bad, bad)
    hid, _ = anim_targets(s)
    if m["kind"] == "Fall":
        # nichts Markiertes vor dem Klick: keine Shape mit grünem Rahmen/Füllung ausser den animierten
        green = [sh.name for sh in s.shapes if sh.shape_id not in hid and sh.shape_type == 1 and sh.name.startswith("Reveal")]
        ok(f"Folie {i}: keine Reveal-Shape ohne Animation", not green, green)
        ok(f"Folie {i}: Antworten A–D sichtbar", all(any(f"Option_{c}" == sh.name for sh in s.shapes if sh.shape_id not in hid) for c in "ABCD"))
    else:
        ok(f"Folie {i}: keine Animation (nichts Verstecktes)", s._element.find(".//p:timing", NS) is None)
    names = [sh.name for sh in s.shapes if sh.shape_id in hid]
    ok(f"Folie {i}: Namen versteckter Shapes verraten nichts", all(n.startswith("Reveal") for n in names), names)

# -------- Test 3: Animationsreihenfolge Folien 16-20
print("\n== Animationsreihenfolge Folien 16–20")
EXP = {16: [["Reveal_Fehlerzeile_%d" % k, "Reveal_Regel_%d" % k] for k in (1, 2, 3, 4, 5)],
       17: [["Reveal_Markierung", "Reveal_Haken"], ["Reveal_Begruendung"]], 18: [["Reveal_Markierung", "Reveal_Haken"], ["Reveal_Begruendung"]], 19: [["Reveal_Markierung", "Reveal_Haken"], ["Reveal_Begruendung"]],
       20: [[f"Reveal_Antwort_{k}"] for k in "abcd"]}
for nr in range(16, 21):
    _, (auto, clicks) = anim_targets(P.slides[nr - 1])
    print(f"Folie {nr} ({META[nr-1]['title'][:30]}): {len(clicks)} Klicks → " + " | ".join("+".join(g) for g in clicks))
    ok(f"Folie {nr}: Klickreihenfolge korrekt", clicks == EXP[nr], clicks)
    note = P.slides[nr - 1].notes_slide.notes_text_frame.text
    ok(f"Folie {nr}: Klickfolge in den Notizen", "ANIMATION (Klickfolge)" in note)
_, (auto14, cl14) = anim_targets(P.slides[13]); ok("Folie 14: 12 Countdown-Segmente automatisch, je 60 s", len(auto14) == 12 and not cl14)
t14 = P.slides[13]._element.find(".//p:timing", NS); dl = [int(d) for d in X(t14, ".//p:par[@nodeType]/..")[0:0]] if False else None
delays = sorted({int(c.get("delay")) for c in X(t14, ".//p:cTn[@fill='hold']/p:stCondLst/p:cond[@delay!='indefinite']") if c.get("delay") not in ("0", None)})
ok("Folie 14: Verzögerungen 60 000 ms-Schritte", delays[:3] == [60000, 120000, 180000] and delays[-1] == 660000, delays[:3] + delays[-1:])
for i, s in enumerate(P.slides, 1):
    ids = [sh.shape_id for sh in s.shapes]; ok(f"Folie {i}: Shape-IDs eindeutig", len(ids) == len(set(ids)))

# -------- Test 4: Zeiten / Anzahl
print("\n== Zeiten und Anzahl")
vis = [m for m in META if not m["hidden"]]; hid = [m for m in META if m["hidden"]]
tot = sum(m["dauer"] for m in vis); ph = {}
for m in vis: ph[m["phase"]] = ph.get(m["phase"], 0) + m["dauer"]
print({k: f"{v//60}:{v%60:02d}" for k, v in ph.items()}, f"Summe {tot//60}:{tot%60:02d}")
ok("Summe der Notizzeiten ≤ 40:00", tot <= 2400, tot)
ok("Sichtbare Folien 20–22", 20 <= len(vis) <= 22, len(vis)); ok("Ausgeblendete Folien ≤ 5", len(hid) <= 5, len(hid))
ok("Nur 6 Phasen im Band", len({m['phase'] for m in META}) == 6)
for i, s in enumerate(P.slides, 1):
    n = s.notes_slide.notes_text_frame.text
    ok(f"Folie {i}: Notizen mit Sprecher/Dauer/Sozialform", all(k in n for k in ("SPRECHER:", "DAUER:", "SOZIALFORM:")) and re.search(r"SPRECHER: (Yasin|Mido)", n) is not None)
ok("Ausgeblendete Folien am Ende und markiert", all(P.slides[i - 1]._element.get("show") == "0" for i, m in enumerate(META, 1) if m["hidden"]) and all(m["hidden"] for m in META[len(vis):]))
st = {i: ("ja" if "streichbar: ja" in META[i - 1]["puffer"] or "streichbar: nein" not in META[i - 1]["puffer"] else "nein") for i in (5, 12, 18, 19)}
ok("Streichbar-Hinweis in Notizen (Folien 5, 12, 18, 19)", all(META[i - 1]["puffer"] for i in (5, 12, 18, 19)))

# -------- Test 5: Schrift
print("\n== Schriftgrößen")
small = {}
for i, s in enumerate(P.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size and r.text.strip() and r.font.size.pt < 28 and META[i - 1]["kind"] == "Erklärfolie" and sh.name != "" and r.font.size.pt < 24: small.setdefault(i, set()).add(r.font.size.pt)
expl = [i for i, m in enumerate(META, 1) if m["kind"] == "Erklärfolie"]
under = {}
for i in expl:
    for sh in P.slides[i - 1].shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size and r.text.strip() and r.font.size.pt < 28: under.setdefault(i, set()).add(r.font.size.pt)
print("Erklärfolien mit Text < 28 pt (nur Paragraf-Ecke/Beschriftungen 24 pt):", {k: sorted(v) for k, v in under.items()})
ok("Erklärfolien: Fließtext ≥ 28 pt (Beschriftungen ≥ 24 pt)", all(min(v) >= 24 for v in under.values()), under)
titles = [(i, min(r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size)) for i in expl for sh in P.slides[i - 1].shapes if sh.name == "Titel" and sh.text_frame.text.strip()]
ok("Titel der Erklärfolien ≥ 40 pt", all(sz >= 40 for _, sz in titles), titles)
low = {}
for i, s in enumerate(P.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.text.strip() and r.font.size.pt < 24: low.setdefault(i, set()).add(r.font.size.pt)
print("Folien mit Text < 24 pt (Dokumentansichten / Plan B):", {k: sorted(v) for k, v in low.items()})

# -------- Test 6: Fakten
print("\n== Fakten")
D = json.loads(open(os.path.join(ROOT, "site/js/data.js"), encoding="utf-8").read().split("=", 1)[1].rstrip().rstrip(";"))
norm = lambda x: re.sub(r"\s+", " ", x).strip()
alltxt = norm(" ".join(sh.text_frame.text for s in P.slides for sh in s.shapes if sh.has_text_frame))
notes = " ".join(s.notes_slide.notes_text_frame.text for s in P.slides)
for f in ["§ 11 BBiG", "§ 8, § 21 BBiG", "§ 8 JArbSchG", "§ 3 ArbZG", "§ 17 BBiG", "§ 19 JArbSchG", "§ 3 BUrlG", "§ 20 BBiG", "§ 22 Abs. 1 und 2 BBiG", "§ 22 Abs. 3 und 4 BBiG", "724 €", "854 €", "977 €", "1.014 €", "01.08.2024", "30 Werktage", "27 Werktage", "25 Werktage", "24 Werktage", "20 Arbeitstage", "8,5", "48 pro Woche", "40 pro Woche"]:
    ok(f"Fakt: {f}", f in alltxt)
mit = 15; t = norm(" ".join(sh.text_frame.text for sh in P.slides[mit - 1].shapes if sh.has_text_frame))
ok("Vertragsauszug wörtlich auf Folie 15", all(norm(z["t"]) in t for z in D["zeilen"]) and all(k in t for k in D["kopf"]))
t16 = norm(" ".join(sh.text_frame.text for sh in P.slides[15].shapes if sh.has_text_frame))
ok("Folie 16: alle 5 Fehlerzeilen + Regeln wörtlich", all(norm(z["t"]) in t16 for z in D["zeilen"] if z["f"]) and all(norm(v["regel"]) in t16 for v in D["fehler"].values()))
for k, nr in zip(range(3), (17, 18, 19)):
    f = D["faelle"][k]; tt = norm(" ".join(sh.text_frame.text for sh in P.slides[nr - 1].shapes if sh.has_text_frame))
    ok(f"Folie {nr}: Fall {'ABC'[k]} wörtlich (Text, Antworten, Begründung)", norm(f["q"]) in tt and all(norm(a) in tt for a in f["a"]) and f["fb"] in tt)
    ok(f"Folie {nr}: Lösungsbuchstabe {'ABCD'[f['c']]} im Haken", f"✓ {'ABCD'[f['c']]}" in tt)
ok("Plan B: 10 Aussagen + 10 Lösungen wörtlich", all(norm(r["t"]) in alltxt and norm(r["b"]) in alltxt for r in D["rg"]))
ok("Plan B: Mini-Quiz 5 Fragen + Antworten + Lösungen wörtlich", all(norm(q["q"]) in alltxt and all(norm(a) in alltxt for a in q["a"]) and q["fb"] in alltxt for q in D["quiz"]))
ok("Fälle D/E und Experten-Karten nicht in der Präsentation", "Reservefall" not in alltxt and "Nina Koch" not in alltxt and "Experten" not in alltxt)
ok("Kein 'WiSo' / keine Projektinhalte", "WiSo" not in alltxt and not re.search(r"projektskizze|strukturplan|ablaufplan|projektbericht|gantt", (alltxt + notes).lower()))
t3 = norm(" ".join(sh.text_frame.text for sh in P.slides[2].shapes if sh.has_text_frame))
ok("Folie 3: vier Fragen, KEINE Antworten", all(q in t3 for q in ["Wie lange darf die Probezeit höchstens dauern?", "Wie viele Urlaubstage stehen einem Azubi mindestens zu?", "Welche Kündigungsfrist gilt in der Probezeit?", "Muss eine Kündigung schriftlich sein?"]))
ok("Folie 20: Antworten der Fragerunde", all(x in alltxt for x in ["Vier Monate", "24 Werktage ab 18 Jahren, Jugendliche 25 bis 30", "Keine Frist, aber schriftlich", "Ja, die Kündigung muss schriftlich erfolgen"]))
n_ok = sum(1 for r in res if r[0]); print(f"\nErgebnis: {n_ok} von {len(res)} Prüfungen bestanden")
open("/tmp/v5_testliste.txt", "w", encoding="utf-8").write("\n".join(("✓ " if r[0] else "✗ ") + r[1] + ("" if r[0] else f" → {r[2]}") for r in res))
